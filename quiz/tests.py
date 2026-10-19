import datetime
import io
import json

from django.test import TestCase, override_settings
from django.urls import reverse

from . import excel
from .models import Answer, Attempt, Choice, Classe, Exam, Question, Subject, User


class QuizFlowTests(TestCase):
    def setUp(self):
        self.c1 = Classe.objects.get(code="1AS")
        self.c7 = Classe.objects.get(code="7SN")
        self.math = Subject.objects.get(name="الرياضيات")
        self.physique = Subject.objects.get(name="الفيزياء والكيمياء")
        self.teacher = User.objects.create_user("prof", password="x", role=User.TEACHER)
        self.teacher.teaching_classes.set([self.c1, self.c7])
        self.teacher.teaching_subjects.set([self.math, self.physique])
        self.other_teacher = User.objects.create_user("prof2", password="x", role=User.TEACHER)
        self.s1 = User.objects.create_user("s1", password="x", role=User.STUDENT, classe=self.c1)
        self.s7 = User.objects.create_user("s7", password="x", role=User.STUDENT, classe=self.c7)
        self.exam = Exam.objects.create(
            title="T", classe=self.c1, subject=self.math, teacher=self.teacher, is_published=True,
        )
        for i, (correct, pts) in enumerate([("B", 2), ("A", 3)]):
            q = Question.objects.create(exam=self.exam, text=f"q{i}", correct_answer=correct, points=pts)
            for l in "ABC":
                Choice.objects.create(question=q, letter=l, text=l)

    def test_seeded_classes_and_subjects(self):
        self.assertEqual(Classe.objects.count(), 11)
        self.assertFalse(self.physique.classes.filter(code="1AS").exists())
        self.assertTrue(self.physique.classes.filter(code="7SN").exists())
        self.assertTrue(Subject.objects.get(name="العربية").classes.filter(code="4AS").exists())

    def test_student_only_sees_own_class(self):
        self.client.login(username="s7", password="x")
        r = self.client.get(reverse("student_dashboard"))
        self.assertNotContains(r, "ابدأ الامتحان")
        r = self.client.get(reverse("exam_intro", args=[self.exam.id]))
        self.assertEqual(r.status_code, 403)
        r = self.client.get(reverse("teacher_dashboard"))
        self.assertEqual(r.status_code, 403)

    def test_full_attempt_and_grading(self):
        self.client.login(username="s1", password="x")
        self.client.post(reverse("exam_intro", args=[self.exam.id]))
        q1, q2 = self.exam.questions.all()
        r = self.client.post(
            reverse("save_answer", args=[self.exam.id]),
            data=json.dumps({"question": q1.id, "selected": "B"}), content_type="application/json",
        )
        self.assertTrue(r.json()["ok"])
        # أثناء الامتحان: أي صفحة أخرى تعيده للامتحان
        r = self.client.get(reverse("student_dashboard"))
        self.assertRedirects(r, reverse("take_exam", args=[self.exam.id]))
        self.client.post(reverse("submit_exam", args=[self.exam.id]), {f"q_{q2.id}": "C"})
        a = Attempt.objects.get(student=self.s1)
        self.assertEqual(a.status, Attempt.SUBMITTED)
        self.assertEqual(a.score, 2)
        self.assertEqual(a.total, 5)
        # لا يمكن إعادة الامتحان
        r = self.client.get(reverse("exam_intro", args=[self.exam.id]))
        self.assertRedirects(r, reverse("attempt_result", args=[self.exam.id]))

    @override_settings(EXAM_MAX_VIOLATIONS=2)
    def test_violations_auto_submit(self):
        self.client.login(username="s1", password="x")
        self.client.post(reverse("exam_intro", args=[self.exam.id]))
        url = reverse("report_violation", args=[self.exam.id])
        self.assertFalse(self.client.post(url).json()["submitted"])
        self.assertTrue(self.client.post(url).json()["submitted"])
        a = Attempt.objects.get(student=self.s1)
        self.assertTrue(a.auto_submitted)

    def test_teacher_cannot_manage_other_exam(self):
        self.client.login(username="prof2", password="x")
        self.assertEqual(self.client.get(reverse("exam_detail", args=[self.exam.id])).status_code, 403)

    def test_exam_subject_must_match_class(self):
        self.client.login(username="prof", password="x")
        r = self.client.post(reverse("exam_create"), {
            "title": "x", "classe": self.c1.id, "subject": self.physique.id, "duration_minutes": 10,
        })
        self.assertContains(r, "لا تُدرَّس")
        self.assertEqual(Exam.objects.count(), 1)

    def test_excel_import_questions(self):
        buf = io.BytesIO()
        excel.questions_template().save(buf)
        buf.seek(0)
        count, errors = excel.import_questions(self.exam, buf)
        self.assertEqual(errors, [])
        self.assertEqual(count, 3)
        last = self.exam.questions.last()
        self.assertEqual(last.correct_answer, "D")  # «د» تحولت إلى D

    def _template_file(self):
        buf = io.BytesIO()
        excel.questions_template().save(buf)
        buf.seek(0)
        buf.name = "q.xlsx"
        return buf

    def test_excel_import_replace_keeps_matching_images(self):
        old = self.exam.questions.first()
        old.text = "كم يساوي 5 + 7 ؟"
        old.image = "questions/fig.png"
        old.save()
        count, errors = excel.import_questions(self.exam, self._template_file(), replace=True)
        self.assertEqual((count, errors), (3, []))
        self.assertEqual(self.exam.questions.count(), 3)  # الأسئلة القديمة حُذفت
        q = self.exam.questions.get(text="كم يساوي 5 + 7 ؟")
        self.assertEqual(q.image.name, "questions/fig.png")
        # استيراد نفس الملف مرة ثانية لا يكرر الأسئلة
        excel.import_questions(self.exam, self._template_file(), replace=True)
        self.assertEqual(self.exam.questions.count(), 3)

    def test_import_view_replace_keeps_past_results(self):
        qs = list(self.exam.questions.all())
        a = Attempt.objects.create(exam=self.exam, student=self.s1, question_order=[q.id for q in qs])
        Answer.objects.create(attempt=a, question=qs[0], selected=qs[0].correct_answer)
        a.submit()
        old_score, old_total = a.score, a.total
        self.assertGreater(old_score, 0)
        self.client.login(username="prof", password="x")
        url = reverse("exam_import", args=[self.exam.id])
        with self.settings(BASE_DIR=self._tmp()):
            self.client.post(url, {"file": self._template_file(), "replace": "1"})
        self.assertEqual(self.exam.questions.count(), 3)  # الأسئلة القديمة لا تظهر في الامتحان
        a.refresh_from_db()
        self.assertEqual((a.score, a.total), (old_score, old_total))  # النتيجة محفوظة
        self.assertEqual(a.answers.count(), 1)
        self.assertEqual({q.id for q in a.graded_questions()}, {q.id for q in qs})
        r = self.client.get(reverse("attempt_report", args=[a.id]))
        self.assertEqual(r.status_code, 200)
        # بدون الاستبدال تُضاف الأسئلة
        self.client.post(url, {"file": self._template_file()})
        self.assertEqual(self.exam.questions.count(), 6)

    def test_excel_import_rejects_bad_answer(self):
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.append(excel.QUESTION_HEADERS)
        wb.active.append(["س", "1", "2", "", "", "C", 1])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        count, errors = excel.import_questions(self.exam, buf)
        self.assertEqual(count, 0)
        self.assertTrue(errors)

    def test_excel_import_students(self):
        buf = io.BytesIO()
        excel.students_template().save(buf)
        buf.seek(0)
        created, updated, errors = excel.import_students(buf)
        self.assertEqual((created, updated, errors), (2, 0, []))
        u = User.objects.get(username="fatima.7sn")
        self.assertEqual(u.classe.code, "7SN")
        self.assertTrue(u.check_password("abcd"))

    def test_class_results_and_export(self):
        Attempt.objects.create(exam=self.exam, student=self.s1, status=Attempt.SUBMITTED, score=4, total=5)
        self.client.login(username="prof", password="x")
        r = self.client.get(reverse("class_results"), {"classe": self.c1.id})
        self.assertContains(r, "s1")
        r = self.client.get(reverse("class_results"), {"classe": self.c1.id, "export": "xlsx"})
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])

    def test_teacher_adds_question_with_image(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from PIL import Image
        buf = io.BytesIO()
        Image.new("RGB", (10, 10), "red").save(buf, "PNG")
        img = SimpleUploadedFile("fig.png", buf.getvalue(), content_type="image/png")
        self.client.login(username="prof", password="x")
        with self.settings(MEDIA_ROOT=self._tmp()):
            r = self.client.post(reverse("question_add", args=[self.exam.id]), {
                "text": "مع صورة", "image": img, "points": "1.5", "correct_answer": "C",
                "choice_A": "a", "choice_B": "b", "choice_C": "c", "choice_D": "",
            })
            self.assertEqual(r.status_code, 302)
            q = self.exam.questions.get(text="مع صورة")
            self.assertTrue(q.image.name.startswith("questions/"))
            self.assertEqual(q.choices.count(), 3)

    def _tmp(self):
        import tempfile
        d = tempfile.mkdtemp()
        self.addCleanup(__import__("shutil").rmtree, d, True)
        return d


class AttemptReportTests(TestCase):
    """تقرير التلميذ: المستوى، مؤشرات العشوائية، المحاور."""

    def setUp(self):
        from datetime import timedelta
        from django.utils import timezone
        from .models import Answer, Attempt, Choice, Classe, Exam, Question, Subject, User
        self.tz, self.td = timezone, timedelta
        self.Answer, self.Attempt = Answer, Attempt
        classe = Classe.objects.get(code="4AS")
        subject = Subject.objects.get(name="الرياضيات")
        self.teacher = User.objects.create_user("t_rep", password="x", role=User.TEACHER)
        self.teacher.teaching_classes.add(classe)
        self.teacher.teaching_subjects.add(subject)
        self.exam = Exam.objects.create(title="امتحان القبول في مدارس المعارف الحرة - الرياضيات", classe=classe,
                                        subject=subject, teacher=self.teacher, is_published=True)
        self.qs = []
        for i in range(40):
            lvl = "E" if i < 18 else ("M" if i < 36 else "H")
            text = f"Calcule : {i} + 3/4 =" if i % 2 else f"Le triangle ABC numéro {i} est :"
            q = Question.objects.create(exam=self.exam, text=text, correct_answer="ABCD"[i % 4],
                                        points={"E": 1, "M": 2, "H": 3}[lvl], level=lvl, order=i)
            Choice.objects.bulk_create([Choice(question=q, letter=l, text=f"{l}{i}") for l in "ABCD"])
            self.qs.append(q)

    def _attempt(self, username, pick, seconds=60):
        from .models import User
        s = User.objects.create_user(username, password="x", role=User.STUDENT, classe=self.exam.classe)
        start = self.tz.now() - self.td(minutes=90)
        a = self.Attempt.objects.create(exam=self.exam, student=s, question_order=[q.id for q in self.qs])
        self.Attempt.objects.filter(pk=a.pk).update(started_at=start)
        a.refresh_from_db()
        for i, q in enumerate(self.qs):
            self.Answer.objects.create(attempt=a, question=q, selected=pick(i, q),
                                       answered_at=start + self.td(seconds=seconds * (i + 1)))
        a.submit()
        return a

    def test_good_student(self):
        from .analysis import analyze
        a = self._attempt("good", lambda i, q: q.correct_answer if i not in (20, 30, 38) else "A")
        r = analyze(a)
        self.assertIn(r["grade"], ("جيد جدًا", "جيد"))
        self.assertEqual(r["verdict_css"], "ok")
        self.assertTrue(r["by_topic"])

    def test_random_student(self):
        from .analysis import analyze
        a = self._attempt("rnd", lambda i, q: "A", seconds=3)
        r = analyze(a)
        self.assertEqual(r["verdict_css"], "ko")
        self.assertTrue(any("الحرف A" in x for x in r["indicators"]))
        self.assertTrue(any("بسرعة" in x for x in r["indicators"]))

    def test_report_page(self):
        a = self._attempt("pg", lambda i, q: q.correct_answer if i % 3 else "B")
        self.client.login(username="t_rep", password="x")
        res = self.client.get(f"/teacher/attempts/{a.id}/report/")
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, "التوصيات")
        res = self.client.get(f"/teacher/exams/{self.exam.id}/results/")
        self.assertContains(res, "تقرير")
        res = self.client.get(f"/teacher/attempts/{a.id}/report/?lang=fr")
        self.assertContains(res, "Recommandations")
        self.assertContains(res, 'dir="ltr"')
        self.assertContains(res, "Imprimer en français")

    def test_french_report_texts(self):
        from .analysis import analyze
        a = self._attempt("frr", lambda i, q: "A", seconds=3)
        r = analyze(a, lang="fr")
        self.assertEqual(r["lang"], "fr")
        self.assertEqual(r["verdict"], "Réponses probablement données au hasard")
        self.assertTrue(all("ال" not in x for x in r["recs"]))


class AttendanceTests(TestCase):
    def setUp(self):
        from .models import Section
        self.Section = Section
        self.s1 = Section.objects.get(code="1AS1")
        self.s2 = Section.objects.get(code="7SN6")
        self.admin = User.objects.create_user("boss", password="x", role=User.ADMIN)
        self.sup = User.objects.create_user("sup", password="x", role=User.SUPERVISOR)
        self.sup.supervised_sections.set([self.s1])
        self.st = User.objects.create_user("1001", password="x", role=User.STUDENT, section=self.s1, full_name="أحمد")
        self.other = User.objects.create_user("2002", password="x", role=User.STUDENT, section=self.s2)

    def _xlsx(self, rows):
        from openpyxl import Workbook
        wb = Workbook()
        for r in rows:
            wb.active.append(r)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        buf.name = "list.xlsx"
        return buf

    def test_sections_seeded_and_linked_to_level(self):
        self.assertEqual(self.Section.objects.count(), 24)
        self.assertEqual(self.st.classe.code, "1AS")
        self.assertEqual(self.s2.classe.code, "7SN")

    def test_supervisor_only_sees_assigned_sections(self):
        self.client.login(username="sup", password="x")
        self.assertRedirects(self.client.get("/"), reverse("attendance_home"))
        r = self.client.get(reverse("attendance_home"))
        self.assertContains(r, "1AS1")
        self.assertNotContains(r, "7SN6")
        self.assertEqual(self.client.get(reverse("section_roll", args=[self.s2.id])).status_code, 403)
        self.assertEqual(self.client.get(reverse("student_file", args=[self.other.id])).status_code, 403)
        self.assertEqual(self.client.get(reverse("teacher_dashboard")).status_code, 403)
        self.assertEqual(self.client.get(reverse("user_list")).status_code, 403)

    def test_phone2_is_whatsapp_link(self):
        self.st.guardian_phone2 = "46 00 00 01"
        self.st.save()
        self.assertEqual(self.st.wa_link, "https://wa.me/22246000001")
        self.client.login(username="sup", password="x")
        r = self.client.get(reverse("student_file", args=[self.st.id]))
        self.assertContains(r, "https://wa.me/22246000001")

    def test_students_sorted_by_matricule(self):
        from .models import STUDENT_ORDER
        for mat in ("10", "9", "100"):
            User.objects.create_user(f"m{mat}", password="x", role=User.STUDENT, section=self.s1, matricule=mat, full_name="ا")
        self.st.matricule = "2"
        self.st.save()
        self.client.login(username="sup", password="x")
        r = self.client.get(reverse("section_roll", args=[self.s1.id]))
        self.assertEqual([x["student"].matricule for x in r.context["rows"]], ["2", "9", "10", "100"])

    def test_roll_call_notes_and_student_file(self):
        from .models import StudentAttendance, StudentNote
        self.client.login(username="sup", password="x")
        url = reverse("section_roll", args=[self.s1.id])
        r = self.client.post(url, {"date": "2026-10-05", "session": 1, f"s_{self.st.id}": "A", f"n_{self.st.id}": "لم يحضر الواجب"})
        self.assertEqual(r.status_code, 302)
        self.client.post(url, {"date": "2026-10-06", "session": 2, f"s_{self.st.id}": "L"})
        self.client.post(url, {"date": "2026-10-06", "session": 2, f"s_{self.st.id}": "L"})  # إعادة الحفظ لا تكرر
        self.client.post(url, {"date": "2026-10-06", "session": 3, f"s_{self.st.id}": "A"})
        self.assertEqual(StudentAttendance.objects.filter(student=self.st).count(), 3)
        r = self.client.get(reverse("section_tools", args=[self.s1.id]) + "?date=2026-10-06&session=1")
        self.assertEqual([len(x["late"]) for x in r.context["summary"]], [0, 1, 0])
        self.assertEqual(StudentNote.objects.get().text, "لم يحضر الواجب")
        r = self.client.get(reverse("section_report", args=[self.s1.id]) + "?from=2026-10-05&to=2026-10-06")
        self.assertContains(r, "غائب")
        self.assertContains(r, "لم يحضر الواجب")
        r = self.client.get(reverse("student_file", args=[self.st.id]))
        self.assertContains(r, "لم يحضر الواجب")
        self.assertEqual(r.context["totals"], {"present": 0, "absent": 2, "late": 1})

    def test_roster_import_keeps_history(self):
        from .models import StudentAttendance
        StudentAttendance.objects.create(student=self.st, date="2026-10-01", status="A")
        self.client.login(username="boss", password="x")
        f = self._xlsx([
            ["رقم قيد الطالب", "القسم", "اسم الطالب", "رقم هاتف الوكيل", "رقم الواتساب", "رقم هاتف الوكيل2"],
            [1001, "1AS2", "أحمد محمد", 22000001, 22000001, None],
            [3003, "7sn6", "مريم", "33000003", "", ""],
        ])
        self.client.post(reverse("roster_import"), {"file": f})
        self.st.refresh_from_db()
        self.assertEqual(self.st.section.code, "1AS2")
        self.assertEqual(self.st.matricule, "1001")
        self.assertEqual(self.st.guardian_phone, "22000001")
        self.assertEqual(self.st.attendance.count(), 1)
        new = User.objects.get(matricule="3003")
        self.assertEqual((new.section.code, new.classe.code, new.username), ("7SN6", "7SN", "3003"))
        self.assertTrue(new.check_password("3003"))
        # إعادة الاستيراد تحدّث ولا تكرر
        self.client.post(reverse("roster_import"), {"file": self._xlsx([[1001, "1AS1", "أحمد", "", "", ""]])})
        self.assertEqual(User.objects.filter(matricule="1001").count(), 1)
        self.assertEqual(self.st.attendance.count(), 1)

    def test_roster_import_detects_columns_in_any_order(self):
        self.client.login(username="boss", password="x")
        f = self._xlsx([
            ["لائحة التلاميذ 2026"],
            ["الرقم", "الإسم الكامل", "الفصل", "الواتساب", "هاتف الولي"],
            [4001, "زيدان محمد العربي", "1AS 1", 46000001, 22000001],
            [4002, "اسماء احمدو", "7sn2", None, 22000002],
        ])
        r = self.client.post(reverse("roster_import"), {"file": f}, follow=True)
        self.assertNotContains(r, "غير موجود")
        u = User.objects.get(matricule="4001")
        self.assertEqual((u.full_name, u.section.code, u.whatsapp, u.guardian_phone),
                         ("زيدان محمد العربي", "1AS1", "46000001", "22000001"))
        self.assertEqual(User.objects.get(matricule="4002").section.code, "7SN2")
        # بدون عناوين وبترتيب مختلف
        f = self._xlsx([[5001, "مريم سيدي", "5MA2", 33000001], [5002, "محمد", "6SN1", 33000002]])
        self.client.post(reverse("roster_import"), {"file": f})
        self.assertEqual(User.objects.get(matricule="5001").section.code, "5MA2")

    def test_roster_import_rejects_unknown_section(self):
        self.client.login(username="boss", password="x")
        r = self.client.post(reverse("roster_import"), {"file": self._xlsx([["رقم القيد", "القسم", "الاسم"], [5, "9XX1", "سالم"]])}, follow=True)
        self.assertContains(r, "غير موجود")
        self.assertFalse(User.objects.filter(matricule="5").exists())

    def test_roster_import_without_section_column_says_so(self):
        self.client.login(username="boss", password="x")
        f = self._xlsx([["رقم القيد", "الاسم"], [7001, "فاطمة وناني باري"], [7002, "محمد سالم"]])
        r = self.client.post(reverse("roster_import"), {"file": f}, follow=True)
        self.assertContains(r, "لم أجد عمود القسم")
        self.assertNotContains(r, "«فاطمة وناني باري» غير موجود")
        self.assertFalse(User.objects.filter(matricule="7001").exists())

    def test_roster_import_section_from_sheet_name_or_title(self):
        from openpyxl import Workbook
        self.client.login(username="boss", password="x")
        wb = Workbook()
        ws = wb.active
        ws.title = "7SN2"
        for row in [["رقم القيد", "الاسم"], [7101, "فاطمة وناني باري"], [7102, "محمد سالم"]]:
            ws.append(row)
        ws2 = wb.create_sheet("ورقة")
        for row in [["لائحة القسم 1AS1"], ["رقم القيد", "الاسم"], [7103, "أحمد"]]:
            ws2.append(row)
        buf = io.BytesIO(); wb.save(buf); buf.seek(0); buf.name = "list.xlsx"
        self.client.post(reverse("roster_import"), {"file": buf}, follow=True)
        self.assertEqual(User.objects.get(matricule="7101").section.code, "7SN2")
        self.assertEqual(User.objects.get(matricule="7101").full_name, "فاطمة وناني باري")
        self.assertEqual(User.objects.get(matricule="7103").section.code, "1AS1")

    def test_roster_import_mislabelled_header_detected_by_content(self):
        self.client.login(username="boss", password="x")
        # عنوان «القسم» فوق عمود الأسماء، والقسم الحقيقي في عمود بلا عنوان واضح
        f = self._xlsx([["رقم القيد", "القسم", "ملاحظة"], [7201, "فاطمة وناني باري", "7SN2"], [7202, "محمد سالم", "7SN2"]])
        self.client.post(reverse("roster_import"), {"file": f}, follow=True)
        u = User.objects.get(matricule="7201")
        self.assertEqual(u.section.code, "7SN2")
        self.assertEqual(u.full_name, "فاطمة وناني باري")

    def test_users_page_import_accepts_roster_and_section_codes(self):
        self.client.login(username="boss", password="x")
        # ملف لائحة التلاميذ يُرفع من صفحة المستخدمين: كان يعطي خطأ 500
        f = self._xlsx([["رقم قيد الطالب", "القسم", "اسم الطالب", "هاتف الوكيل"], [6001, "6AF", "سالم", 22000006]])
        r = self.client.post(reverse("students_import"), {"file": f}, follow=True)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(User.objects.get(matricule="6001").section.code, "6AF")
        # الصيغة القديمة مع رمز فصل مثل 1AS1، وقسم غير معروف لا يسبب خطأ 500
        f = self._xlsx([["اسم المستخدم", "الاسم الكامل", "كلمة المرور", "القسم"], ["ali", "علي", "1234", "1AS1"]])
        self.client.post(reverse("students_import"), {"file": f})
        self.assertEqual(User.objects.get(username="ali").section.code, "1AS1")
        f = self._xlsx([["اسم المستخدم", "الاسم الكامل", "كلمة المرور", "القسم"], ["bad", "x", "1", None]])
        r = self.client.post(reverse("students_import"), {"file": f}, follow=True)
        self.assertContains(r, "غير موجود")

    def test_reimport_does_not_log_students_out(self):
        student = self.client_class()
        student.login(username="1001", password="x")
        self.assertEqual(student.get(reverse("home")).wsgi_request.user.username, "1001")
        self.client.login(username="boss", password="x")
        f = self._xlsx([["اسم المستخدم", "الاسم الكامل", "كلمة المرور", "القسم"], ["1001", "أحمد", "x", "1AS1"]])
        self.client.post(reverse("students_import"), {"file": f})
        f = self._xlsx([["رقم قيد الطالب", "القسم", "اسم الطالب"], [1001, "1AS2", "أحمد"]])
        self.client.post(reverse("roster_import"), {"file": f})
        self.assertEqual(student.get(reverse("home")).wsgi_request.user.username, "1001")

    def test_teachers_import_and_monthly(self):
        from .models import Teacher, TeacherAttendance
        self.client.login(username="boss", password="x")
        self.client.post(reverse("teachers_import"), {"file": self._xlsx([["اسم الأستاذ", "المادة"], ["محمد", "الرياضيات"], ["Aicha", "الفرنسية"]])})
        self.assertEqual(Teacher.objects.count(), 2)
        t = Teacher.objects.get(name="محمد")
        self.client.logout()
        self.client.login(username="sup", password="x")
        self.client.post(reverse("teacher_roll"), {"date": "2026-10-05", "session": 1, f"s_{t.id}": "A", f"c_{t.id}": self.s1.id})
        self.client.post(reverse("teacher_roll"), {"date": "2026-10-06", "session": 1, f"s_{t.id}": "P", f"c_{t.id}": self.s1.id})
        self.client.post(reverse("teacher_roll"), {"date": "2026-10-06", "session": 3, f"s_{t.id}": "P", f"c_{t.id}": self.s2.id})
        self.assertEqual(
            list(TeacherAttendance.objects.filter(teacher=t, date="2026-10-06").values_list("session", "section__code")),
            [(1, "1AS1"), (3, "7SN6")],
        )
        # «لا حصة» يلغي التسجيل، واليوم التالي يبدأ فارغاً
        self.client.post(reverse("teacher_roll"), {"date": "2026-10-06", "session": 3, f"s_{t.id}": ""})
        self.assertEqual(TeacherAttendance.objects.filter(teacher=t, date="2026-10-06").count(), 1)
        r = self.client.get(reverse("teacher_roll") + "?date=2026-10-07&session=1")
        row = [x for x in r.context["rows"] if x["teacher"].id == t.id][0]
        self.assertEqual((row["status"], row["section"]), ("", None))
        self.client.logout()
        self.client.login(username="boss", password="x")
        self.client.post(reverse("teachers_import"), {"file": self._xlsx([["محمد", "الفيزياء"]])})
        self.assertEqual(TeacherAttendance.objects.filter(teacher=t).count(), 2)
        self.client.logout()
        self.client.login(username="sup", password="x")
        r = self.client.get(reverse("teacher_monthly") + "?month=2026-10")
        row = [x for x in r.context["teachers"] if x.id == t.id][0]
        self.assertEqual((row.n_present, row.n_absent, row.subject), (1, 1, "الفيزياء"))
        self.assertEqual(row.month_sections, ["1AS1"])
        self.assertEqual(self.client.get(reverse("teacher_file", args=[t.id]) + "?month=2026-10").status_code, 200)
        # المراقب لا يستورد
        self.assertEqual(self.client.post(reverse("teachers_import")).status_code, 403)

    def test_students_list_and_export(self):
        self.st.matricule, self.st.guardian_phone = "1001", "22000001"
        self.st.save()
        self.client.login(username="sup", password="x")
        r = self.client.get(reverse("students_list"))
        self.assertContains(r, "22000001")
        self.assertNotContains(r, "2002")  # تلميذ قسم لا يراقبه
        r = self.client.get(reverse("students_list") + "?section=1AS1&export=xlsx")
        self.assertEqual(r.status_code, 200)
        from openpyxl import load_workbook
        rows = list(load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))
        self.assertEqual(rows[1][1:5], ("1001", "أحمد", "1AS1", "22000001"))
        self.client.logout()
        self.client.login(username="boss", password="x")
        self.assertContains(self.client.get(reverse("students_list")), "2002")

    def test_section_tabs_students_and_absentees(self):
        self.st.guardian_phone = "22000001"
        self.st.save()
        other = User.objects.create_user("1009", password="x", role=User.STUDENT, section=self.s1, full_name="حاضر")
        self.client.login(username="sup", password="x")
        self.client.post(reverse("section_roll", args=[self.s1.id]),
                         {"date": "2026-10-05", "session": 2, f"s_{self.st.id}": "A", f"s_{other.id}": "P"})
        r = self.client.get(reverse("section_absentees", args=[self.s1.id]) + "?date=2026-10-05&session=2")
        self.assertContains(r, "22000001")
        self.assertNotContains(r, ">حاضر</a>")
        self.assertContains(r, "لائحة القسم")
        r = self.client.get(reverse("section_absentees", args=[self.s1.id]) + "?date=2026-10-05&session=1")
        self.assertContains(r, "لم يُسجَّل غياب هذه الحصة")
        r = self.client.get(reverse("section_absentees", args=[self.s1.id]) + "?date=2026-10-05&session=2&export=xlsx")
        from openpyxl import load_workbook
        rows = list(load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))
        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[2][4], "22000001")
        r = self.client.get(reverse("section_students", args=[self.s1.id]))
        self.assertContains(r, "حاضر")
        self.assertEqual(self.client.get(reverse("section_students", args=[self.s1.id]) + "?export=xlsx").status_code, 200)
        self.assertEqual(self.client.get(reverse("section_students", args=[self.s2.id])).status_code, 403)

    def test_roll_sheet_print_and_excel(self):
        self.client.login(username="sup", password="x")
        url = reverse("section_sheet", args=[self.s1.id])
        r = self.client.get(url + "?date=2026-10-05&session=2")
        self.assertContains(r, "أحمد")
        self.assertContains(r, "الحصة 2")
        r = self.client.get(url + "?date=2026-10-05&session=2&export=xlsx")
        from openpyxl import load_workbook
        rows = list(load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))
        self.assertEqual(rows[4][:3], ("#", "رقم القيد", "اسم التلميذ"))
        self.assertEqual(rows[5][2], "أحمد")
        self.assertEqual(self.client.get(reverse("section_sheet", args=[self.s2.id])).status_code, 403)

    def test_admin_creates_supervisor(self):
        self.client.login(username="boss", password="x")
        r = self.client.post(reverse("user_create"), {
            "username": "mourakib", "full_name": "م", "role": "supervisor", "password": "p",
            "supervised_sections": [self.s1.id, self.s2.id], "is_active": "on",
        })
        self.assertEqual(r.status_code, 302)
        u = User.objects.get(username="mourakib")
        self.assertEqual(u.supervised_sections.count(), 2)
        for name in ["user_create", "attendance_home", "teacher_roll", "teacher_monthly"]:
            self.assertEqual(self.client.get(reverse(name)).status_code, 200)


class LibraryTests(TestCase):
    def setUp(self):
        from .models import Section
        self.sec = Section.objects.get(code="7SN6")
        self.sec2 = Section.objects.get(code="7SN1")
        self.lib = User.objects.create_user("lib", password="x", role=User.LIBRARIAN)
        self.a = User.objects.create_user("a1", password="x", role=User.STUDENT, section=self.sec, full_name="أ")
        self.b = User.objects.create_user("b1", password="x", role=User.STUDENT, section=self.sec, full_name="ب")
        self.c = User.objects.create_user("c1", password="x", role=User.STUDENT, section=self.sec2, full_name="ج")

    def test_level_stock_shared_by_sections(self):
        from .models import BookDelivery
        url = reverse("library_section", args=[self.sec.id])
        self.client.login(username="lib", password="x")
        self.assertRedirects(self.client.get(reverse("home")), reverse("library_home"), fetch_redirect_response=False)
        home = self.client.get(reverse("library_home"))
        self.assertNotContains(home, "1AS1")  # من 3AS فقط
        self.client.post(reverse("library_level", args=[self.sec.classe_id]), {"stock_math_1": "30", "stock_math_2": "25"})
        # تلميذ يستلم النسخة 1، والمتأخر يستلم النسختين معاً، وتلميذ من فصل آخر لنفس المستوى
        self.client.post(url, {f"d_{self.a.id}_math_1": "1", f"d_{self.b.id}_math_1": "1", f"d_{self.b.id}_math_2": "1"})
        self.client.post(reverse("library_section", args=[self.sec2.id]), {f"d_{self.c.id}_math_1": "1"})
        self.assertEqual(BookDelivery.objects.count(), 4)
        r = self.client.get(url)
        math1, math2 = r.context["summary"][0]["parts"]
        self.assertEqual((math1["stock"], math1["given"], math1["left"], math1["in_section"]), (30, 3, 27, 2))
        self.assertEqual(math2["left"], 24)
        # إلغاء علامة يلغي التسليم
        self.client.post(url, {f"d_{self.b.id}_math_1": "1", f"d_{self.b.id}_math_2": "1"})
        self.assertFalse(BookDelivery.objects.filter(student=self.a).exists())
        self.assertEqual(self.client.get(url + "?export=xlsx").status_code, 200)

    def test_others_cannot_open_library(self):
        self.client.login(username="a1", password="x")
        self.assertEqual(self.client.get(reverse("library_home")).status_code, 403)


class ActivityLogTests(TestCase):
    def setUp(self):
        from .models import ActivityLog  # noqa: F401
        self.c1 = Classe.objects.get(code="1AS")
        self.math = Subject.objects.get(name="الرياضيات")
        self.admin = User.objects.create_user("boss", password="x", role=User.ADMIN, full_name="المدير")
        self.teacher = User.objects.create_user("prof", password="x", role=User.TEACHER, full_name="الأستاذ")
        self.teacher.teaching_classes.set([self.c1])
        self.teacher.teaching_subjects.set([self.math])
        self.s1 = User.objects.create_user("s1", password="x", role=User.STUDENT, classe=self.c1)
        self.exam = Exam.objects.create(title="امتحان تجريبي", classe=self.c1, subject=self.math,
                                        teacher=self.teacher, is_published=True)
        q = Question.objects.create(exam=self.exam, text="q", correct_answer="A", points=1)
        for l in "AB":
            Choice.objects.create(question=q, letter=l, text=l)

    def logs(self):
        from .models import ActivityLog
        return ActivityLog.objects.order_by("created_at")

    def test_login_logout_and_failed_login_are_logged(self):
        self.client.post(reverse("login"), {"username": "prof", "password": "wrong"})
        self.client.post(reverse("login"), {"username": "prof", "password": "x"})
        self.assertIsNotNone(User.objects.get(pk=self.teacher.pk).last_seen)
        self.client.post(reverse("logout"))
        kinds = list(self.logs().values_list("kind", "username"))
        self.assertEqual(kinds, [("failed", "الأستاذ"), ("login", "الأستاذ"), ("logout", "الأستاذ")])
        self.assertIsNone(User.objects.get(pk=self.teacher.pk).last_seen)

    def test_changes_are_logged_with_target_even_after_delete(self):
        self.client.login(username="prof", password="x")
        self.client.post(reverse("exam_toggle_publish", args=[self.exam.id]))
        self.client.post(reverse("exam_delete", args=[self.exam.id]))
        rows = list(self.logs().exclude(kind="login").values_list("kind", "action", "target"))
        self.assertEqual(rows[0], ("edit", "إخفاء امتحان", "امتحان تجريبي (1AS)"))
        self.assertEqual(rows[1], ("delete", "حذف امتحان", "امتحان تجريبي (1AS)"))

    def test_student_exam_events_but_not_each_answer(self):
        self.client.login(username="s1", password="x")
        self.client.post(reverse("exam_intro", args=[self.exam.id]))
        q = self.exam.questions.first()
        self.client.post(reverse("save_answer", args=[self.exam.id]),
                         data=json.dumps({"question": q.id, "selected": "A"}), content_type="application/json")
        self.client.post(reverse("submit_exam", args=[self.exam.id]))
        actions = list(self.logs().exclude(kind="login").values_list("action", flat=True))
        self.assertEqual(actions, ["بدء امتحان", "تسليم امتحان"])

    def test_report_page_shows_online_users_and_exports(self):
        self.client.login(username="s1", password="x")
        self.client.get(reverse("student_dashboard"))
        self.client.login(username="boss", password="x")
        r = self.client.get(reverse("activity_report"))
        self.assertEqual(r.status_code, 200)
        self.assertIn(self.s1, r.context["online"])
        self.assertContains(r, "تسجيل الدخول")
        r = self.client.get(reverse("activity_report") + "?mode=month")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(len(r.context["per_day"]), 1)
        r = self.client.get(reverse("activity_report") + "?mode=month&export=1")
        self.assertEqual(r["Content-Type"].split(";")[0],
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def test_report_is_admin_only(self):
        self.client.login(username="prof", password="x")
        self.assertEqual(self.client.get(reverse("activity_report")).status_code, 403)

    def test_superuser_is_logged_as_admin_not_student(self):
        User.objects.create_superuser("root", password="x")  # خانة الدور تبقى «تلميذ» افتراضياً
        self.client.post(reverse("login"), {"username": "root", "password": "x"})
        self.assertEqual(self.logs().get(kind="login").role, User.ADMIN)
        r = self.client.get(reverse("activity_report"))
        self.assertContains(r, "<td>مشرف</td>", html=True)


class TimetableTests(TestCase):
    def setUp(self):
        from .models import Section
        self.c1 = Classe.objects.get(code="1AS")
        self.admin = User.objects.create_user("boss", password="x", role=User.ADMIN)
        self.s1 = User.objects.create_user("s1", password="x", role=User.STUDENT, classe=self.c1,
                                           section=Section.objects.get(code="1AS1"))
        self.s2 = User.objects.create_user("s2", password="x", role=User.STUDENT, classe=self.c1,
                                           section=Section.objects.get(code="1AS2"))

    def general_file(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "General"
        ws["D4"] = "PLANNING"
        ws.append([])
        for col, code in zip("CDEF", ["1A1", "1A2", "5C1", "7D9"]):
            ws[f"{col}5"] = code
        ws["J5"] = "6D3"  # ملاحظة جانبية بعد فراغ: ليست فصلاً
        rows = [
            ("Lun", "8-10", "AR3", "M10-M8", "PC(3,5)", "Ti"),
            (None, "10-12", "Fr1", None, "SN1", "CH"),
            (None, None, None, None, None, None),
            ("Mar", "8-10", "M1", "HG1", None, None),
        ]
        for i, r in enumerate(rows, start=6):
            for j, v in enumerate(r, start=1):
                ws.cell(i, j, v)
        for col, code in zip("CDEF", ["1A1", "1A2", "5C1", "7D9"]):
            ws[f"{col}10"] = code
        ws["J10"], ws["K10"] = "المادة", "الاسم"
        legend = [("AR3", "محفوظ  محمد فال"), ("M10", "احمد اسحاق"), ("M8-M6", "حرمه بن حمود"),
                  ("Fr1", "دم مامادو"), ("PC3", "أستاذ الفيزياء")]
        for i, (code, name) in enumerate(legend, start=11):
            ws[f"J{i}"], ws[f"K{i}"] = code, name
        ws["A30"], ws["B30"], ws["C30"] = "Lun", "8-10", "IR9"  # جدول قديم أسفل الورقة
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        buf.name = "emploi.xlsx"
        return buf

    def test_import_splits_general_into_sections(self):
        from .models import TimetableEntry
        self.client.login(username="boss", password="x")
        resp = self.client.post(reverse("timetable_import"), {"file": self.general_file()}, follow=True)
        self.assertContains(resp, "4 قسماً")
        self.assertEqual(sorted(set(TimetableEntry.objects.values_list("section_code", flat=True))),
                         ["1AS1", "1AS2", "5MA1", "7SN9"])
        e = TimetableEntry.objects.get(section_code="1AS1", day=0, slot="8-10")
        self.assertEqual((e.subject, e.teacher), ("اللغة العربية", "محفوظ محمد فال"))
        e = TimetableEntry.objects.get(section_code="1AS2", day=0)
        self.assertEqual((e.subject, e.teacher), ("الرياضيات", "احمد اسحاق / حرمه بن حمود"))
        e = TimetableEntry.objects.get(section_code="5MA1", day=0, slot="8-10")
        self.assertEqual((e.subject, e.teacher), ("الفيزياء والكيمياء", "أستاذ الفيزياء"))
        self.assertEqual(TimetableEntry.objects.get(section_code="7SN9", slot="8-10").subject, "المعلوماتية والتكنولوجيا")
        self.assertEqual(TimetableEntry.objects.get(section_code="7SN9", slot="10-12").subject, "الكيمياء")
        self.assertFalse(TimetableEntry.objects.filter(code="IR9").exists())
        self.assertContains(resp, "7SN9")  # فصل غير موجود في التطبيق
        # إعادة الاستيراد تحلّ محلّ الجدول القديم
        self.client.post(reverse("timetable_import"), {"file": self.general_file()})
        self.assertEqual(TimetableEntry.objects.filter(section_code="1AS1").count(), 3)

    def test_student_sees_only_own_section(self):
        from .timetable import import_timetable
        import_timetable(self.general_file())
        self.client.login(username="s1", password="x")
        resp = self.client.get(reverse("my_timetable"))
        self.assertContains(resp, "اللغة الفرنسية")
        self.assertNotContains(resp, "التاريخ والجغرافيا")
        self.assertNotContains(resp, "محفوظ محمد فال")  # اسم الأستاذ لا يظهر
        self.assertContains(self.client.get(reverse("student_dashboard")), "اللغة الفرنسية")
        self.assertEqual(self.client.get(reverse("timetable_home")).status_code, 403)
        self.client.login(username="s2", password="x")
        self.assertContains(self.client.get(reverse("my_timetable")), "التاريخ والجغرافيا")

    def test_only_admin_imports(self):
        teacher = User.objects.create_user("prof", password="x", role=User.TEACHER)
        self.client.login(username="prof", password="x")
        self.assertEqual(self.client.get(reverse("timetable_home")).status_code, 200)
        resp = self.client.post(reverse("timetable_import"), {"file": self.general_file()})
        self.assertEqual(resp.status_code, 403)
        self.assertIsNotNone(teacher)

    def test_bad_file_keeps_old_timetable(self):
        from openpyxl import Workbook
        from .models import TimetableEntry
        from .timetable import import_timetable
        import_timetable(self.general_file())
        buf = io.BytesIO()
        Workbook().save(buf)
        buf.seek(0)
        buf.name = "x.xlsx"
        self.client.login(username="boss", password="x")
        resp = self.client.post(reverse("timetable_import"), {"file": buf}, follow=True)
        self.assertContains(resp, "لم يتغيّر الجدول السابق")
        self.assertTrue(TimetableEntry.objects.exists())


@override_settings(MEDIA_ROOT="/tmp/claude-0/test-media")
class LessonTests(TestCase):
    def setUp(self):
        from .models import LessonScope, Section
        self.sec = Section.objects.get(code="2AS1")
        self.other = Section.objects.get(code="2AS2")
        self.admin = User.objects.create_user("boss", password="x", role=User.ADMIN)
        self.rec = User.objects.create_user("rec", password="x", role=User.RECORDER, full_name="المسجّل")
        self.scope = LessonScope.objects.create(user=self.rec, section=self.sec, subject="M")
        self.s1 = User.objects.create_user("s1", password="x", role=User.STUDENT, section=self.sec)

    def photo(self, w=1200, h=1600):
        from PIL import Image
        buf = io.BytesIO()
        Image.new("RGB", (w, h), "white").save(buf, "JPEG")
        buf.seek(0)
        buf.name = "p.jpg"
        return buf

    def test_recorder_saves_lesson_and_pdf_grows(self):
        from .models import Lesson
        self.client.login(username="rec", password="x")
        self.assertContains(self.client.get(reverse("lessons_home")), "تصوير درس جديد")
        url = reverse("lesson_new", args=[self.scope.id])
        r = self.client.post(url, {"date": "2026-10-12", "title": "الدرس 1", "pages": [self.photo(), self.photo(3000, 4000)]})
        self.assertRedirects(r, reverse("lessons_home"))
        lesson = Lesson.objects.get()
        self.assertEqual((lesson.subject, lesson.section, lesson.pages.count()), ("M", self.sec, 2))
        self.assertLessEqual(max(lesson.pages.last().image.width, lesson.pages.last().image.height), 1600)
        self.client.post(url, {"date": "2026-10-14", "pages": [self.photo()]})
        pdf = self.client.get(reverse("lesson_pdf", args=[self.sec.id, "M"]))
        self.assertEqual(pdf["Content-Type"], "application/pdf")
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        self.assertIn(b"/Count 3", pdf.content)
        one_day = self.client.get(reverse("lesson_pdf", args=[self.sec.id, "M"]) + "?from=2026-10-14&to=2026-10-14")
        self.assertIn(b"/Count 1", one_day.content)
        self.client.login(username="boss", password="x")
        r = self.client.get(reverse("lesson_board", args=[self.sec.id, "M"]))
        self.assertContains(r, "الدرس 2")
        self.assertEqual(r.content.count(b'class="nb-page"'), 3)
        z = self.client.get(reverse("lessons_zip"))
        import zipfile
        names = zipfile.ZipFile(io.BytesIO(b"".join(z.streaming_content))).namelist()
        self.assertEqual(names, ["2AS1/2AS1_Maths.pdf"])

    def test_recorder_limited_to_assigned_scopes(self):
        from .models import LessonScope
        other_scope = LessonScope.objects.create(user=self.admin, section=self.other, subject="AR")
        self.client.login(username="rec", password="x")
        self.assertEqual(self.client.get(reverse("lesson_new", args=[other_scope.id])).status_code, 404)
        self.assertEqual(self.client.get(reverse("lesson_board", args=[self.other.id, "AR"])).status_code, 403)
        self.assertEqual(self.client.get(reverse("lessons_overview")).status_code, 403)
        r = self.client.post(reverse("lesson_new", args=[self.scope.id]), {"date": "2026-10-12"})
        self.assertContains(r, "صوّر صفحة واحدة على الأقل")

    def test_admin_overview_and_scopes(self):
        from .models import Lesson, LessonScope
        Lesson.objects.create(section=self.sec, subject="M", date=datetime.date(2026, 10, 15),
                              title="المعادلات", uploaded_by=self.rec)
        self.client.login(username="boss", password="x")
        r = self.client.get(reverse("lessons_overview"))
        self.assertContains(r, "المعادلات")
        self.assertContains(r, "15/10/2026")
        self.client.post(reverse("lesson_scopes", args=[self.rec.id]),
                         {"scope": [f"{self.other.id}:AR", f"{self.sec.id}:M"]})
        self.assertEqual(set(LessonScope.objects.filter(user=self.rec).values_list("section__code", "subject")),
                         {("2AS2", "AR"), ("2AS1", "M")})
        self.client.login(username="s1", password="x")
        self.assertEqual(self.client.get(reverse("lessons_overview")).status_code, 403)

    def test_new_recorder_account_goes_to_scopes(self):
        self.client.login(username="boss", password="x")
        r = self.client.post(reverse("user_create"), {"username": "rec2", "full_name": "x", "role": "recorder",
                                                      "password": "pw12345", "is_active": "on"})
        self.assertRedirects(r, reverse("lesson_scopes", args=[User.objects.get(username="rec2").id]))
