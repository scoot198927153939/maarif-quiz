"""الجدول الزمني: استيراد الجدول العام من Excel وعرض جدول كل قسم."""
import re

from django.contrib import messages
from django.db import transaction
from django.shortcuts import redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST
from openpyxl import load_workbook

from .forms import ImportFileForm
from .models import Section, TimetableEntry
from .views import admin_required, role_required, student_required

timetable_viewer_required = role_required("admin", "supervisor", "teacher")

SUBJECTS = {
    "AR": "اللغة العربية",
    "FR": "اللغة الفرنسية",
    "AN": "اللغة الإنجليزية",
    "M": "الرياضيات",
    "PC": "الفيزياء والكيمياء",
    "SN": "العلوم الطبيعية",
    "HG": "التاريخ والجغرافيا",
    "IR": "التربية الإسلامية",
    "IC": "التربية المدنية",
    "PH": "الفلسفة",
    "TI": "المعلوماتية والتكنولوجيا",
}

# 1A1 → 1AS1، 5C1 → 5MA1، 7D1 → 7SN1 (ورموز التطبيق نفسها مقبولة أيضاً)
SECTION_LETTERS = {"A": "AS", "AS": "AS", "C": "MA", "MA": "MA", "D": "SN", "DN": "SN", "SN": "SN"}
SECTION_RE = re.compile(r"^(\d)\s*([A-Za-z]{1,2})\s*(\d{1,2})$")
SLOT_RE = re.compile(r"^\d{1,2}\s*(h|H)?\s*[-–:]\s*\d{1,2}\s*(h|H)?$")
DAY_NAMES = {
    "lun": 0, "lundi": 0, "الاثنين": 0, "الإثنين": 0,
    "mar": 1, "mardi": 1, "الثلاثاء": 1,
    "mer": 2, "mercredi": 2, "الأربعاء": 2, "الاربعاء": 2,
    "jeu": 3, "jeudi": 3, "الخميس": 3,
    "ven": 4, "vendredi": 4, "الجمعة": 4,
    "sam": 5, "samedi": 5, "السبت": 5,
    "dim": 6, "dimanche": 6, "الأحد": 6, "الاحد": 6,
}
ARABIC_RE = re.compile(r"[؀-ۿ]")


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def section_code(value):
    """يحوّل رمز القسم في الملف إلى رمز التطبيق، أو "" إن لم يكن رمز قسم."""
    m = SECTION_RE.match(_text(value).replace(",", "").strip())
    if not m:
        return ""
    letters = SECTION_LETTERS.get(m.group(2).upper())
    return f"{m.group(1)}{letters}{int(m.group(3))}" if letters else ""


def split_codes(value):
    """"M10-M8" → [M10, M8]، "PC(3,5)" → [PC3, PC5]، "M9-8" → [M9, M8]، "Ti" → [TI]."""
    text = re.sub(r"[\s()]", "", _text(value).upper())
    codes, prefix = [], ""
    for part in re.split(r"[-,+/;]", text):
        m = re.match(r"^([A-Z]*)(\d*)$", part)
        if not part or not m:
            continue
        prefix = m.group(1) or prefix
        if prefix:
            codes.append(prefix + m.group(2))
    return codes


def _subject_key(code):
    return re.match(r"^[A-Z]*", code).group(0)


def _find_header(ws):
    """أول صف فيه رموز أقسام متتالية (3 على الأقل). يعيد (رقم الصف، {عمود: رمز})."""
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 40)):
        columns, started = {}, False
        for cell in row:
            code = section_code(cell.value)
            if code:
                started = True
                columns.setdefault(cell.column, code)
            elif started:
                break  # ما بعد أول فراغ ملاحظات جانبية
        if len(columns) >= 3:
            return row[0].row, columns
    return None, {}


def _pick_sheet(wb):
    for ws in wb.worksheets:
        if ws.title.strip().lower() in ("general", "général", "العام", "الجدول العام"):
            return ws
    best, best_n = None, 0
    for ws in wb.worksheets:
        n = len(_find_header(ws)[1])
        if n > best_n:
            best, best_n = ws, n
    return best


def parse_general(file):
    """يقرأ الجدول العام. يعيد (الحصص، رموز الأساتذة، الأخطاء)."""
    try:
        wb = load_workbook(file, data_only=True)
    except Exception:
        return [], {}, ["تعذّرت قراءة الملف. تأكد أنه ملف Excel ‎.xlsx سليم."]
    ws = _pick_sheet(wb)
    header_row, columns = _find_header(ws) if ws else (None, {})
    if not columns:
        return [], {}, ["لم أجد صف رموز الأقسام (مثل 1A1، 5C1، 7D1) في الملف."]

    first_col = min(columns)
    slots, entries, day, last_row = [], [], None, header_row
    for row in ws.iter_rows(min_row=header_row + 1):
        left = [_text(c.value) for c in row[: first_col - 1]]
        if row[0].row > header_row + 1 and sum(1 for c in row if section_code(c.value)) >= 3:
            break  # تكرار صف الرموز أسفل الجدول
        for value in left:
            if value.lower().rstrip(".") in DAY_NAMES:
                day = DAY_NAMES[value.lower().rstrip(".")]
        slot = next((v.replace(" ", "") for v in left if SLOT_RE.match(v)), "")
        if not slot or day is None:
            continue
        last_row = row[0].row
        if slot not in slots:
            slots.append(slot)
        for cell in row:
            code = columns.get(cell.column)
            value = _text(cell.value)
            if code and value:
                entries.append({"section": code, "day": day, "slot": slot,
                                "slot_order": slots.index(slot), "code": value})
    if not entries:
        return [], {}, ["وجدت رموز الأقسام لكن لم أجد أي حصة تحتها (الأيام والتوقيت في أول عمودين)."]

    # دليل الأساتذة على يمين الجدول وأسفله: رمز الأستاذ ثم اسمه في نفس الصف.
    # ينتهي الدليل عند أول 5 صفوف متتالية بلا رمز (ما بعدها جداول قديمة محفوظة).
    name_col = None  # عمود "الاسم" إن وُجد عنوانه
    for row in ws.iter_rows(min_row=header_row, max_row=last_row + 3, min_col=max(columns) + 1):
        for cell in row:
            if _text(cell.value).startswith("الاسم"):
                name_col = name_col or cell.column
    names, shared, empty = {}, {}, 0
    for row in ws.iter_rows(min_row=last_row + 1, min_col=max(columns) + 1):
        values = [_text(c.value) for c in row]
        code_at = next((i for i, v in enumerate(values) if re.match(r"^[A-Za-z]{1,3}\s*\(?\s*\d", v)), None)
        code_cell = values[code_at] if code_at is not None else ""
        codes = split_codes(code_cell)
        if not codes:
            empty += 1
            if empty >= 5:
                break
            continue
        empty = 0
        if name_col:
            name = _text(ws.cell(row[0].row, name_col).value)
        else:
            name = next((v for v in values[code_at + 1:] if ARABIC_RE.search(v)
                         and not re.search(r"\d", v)), "")
        name = " ".join(name.split())
        if not name:
            continue
        if len(codes) == 1:
            names[codes[0]] = name
        else:
            for c in codes:
                shared.setdefault(c, name)
    for c, name in shared.items():
        names.setdefault(c, name)
    return entries, names, []


def import_timetable(file):
    entries, names, errors = parse_general(file)
    if errors:
        return [], errors
    now = timezone.now()
    objs = []
    for e in entries:
        codes = split_codes(e["code"])
        key = _subject_key(codes[0]) if codes else ""
        teachers = []
        for c in codes:
            if names.get(c) and names[c] not in teachers:
                teachers.append(names[c])
        objs.append(TimetableEntry(
            section_code=e["section"], day=e["day"], slot=e["slot"], slot_order=e["slot_order"],
            code=e["code"], subject_key=key, subject=SUBJECTS.get(key, e["code"]),
            teacher=" / ".join(teachers), created_at=now,
        ))
    with transaction.atomic():
        TimetableEntry.objects.all().delete()
        TimetableEntry.objects.bulk_create(objs)
    return sorted({o.section_code for o in objs}, key=_section_sort), []


def _section_sort(code):
    return (code[:1], code[1:-1], int(re.sub(r"\D", "", code[2:]) or 0))


def build_grid(code):
    """جدول قسم: صف لكل يوم وعمود لكل توقيت."""
    entries = list(TimetableEntry.objects.filter(section_code=code))
    if not entries:
        return None
    all_slots = (TimetableEntry.objects.values_list("slot_order", "slot")
                 .order_by("slot_order").distinct())
    slots = []
    for order, label in all_slots:
        if label not in slots:
            slots.append(label)
    days = sorted(set(TimetableEntry.objects.values_list("day", flat=True)))
    cells = {(e.day, e.slot): e for e in entries}
    today = timezone.localdate().weekday()
    rows = [{
        "day": TimetableEntry.DAYS[d], "today": d == today,
        "cells": [cells.get((d, s)) for s in slots],
    } for d in days]
    return {"code": code, "slots": slots, "rows": rows, "count": len(entries),
            "updated": max(e.created_at for e in entries)}


def _codes():
    return sorted(set(TimetableEntry.objects.values_list("section_code", flat=True)), key=_section_sort)


@timetable_viewer_required
def timetable_home(request):
    codes = _codes()
    selected = request.GET.get("s", "")
    show_all = request.GET.get("all") == "1"
    if show_all:
        grids = [build_grid(c) for c in codes]
    elif selected in codes:
        grids = [build_grid(selected)]
    else:
        grids = []
    known = set(Section.objects.values_list("code", flat=True))
    return render(request, "quiz/timetable/home.html", {
        "codes": codes, "selected": selected, "show_all": show_all, "grids": grids,
        "missing": [c for c in codes if c not in known],
        "no_timetable": sorted(known - set(codes), key=_section_sort),
        "import_form": ImportFileForm(),
    })


@admin_required
@require_POST
def timetable_import(request):
    form = ImportFileForm(request.POST, request.FILES)
    if not form.is_valid():
        for e in form.errors.get("file", []):
            messages.error(request, e)
        return redirect("timetable_home")
    codes, errors = import_timetable(form.cleaned_data["file"])
    for e in errors:
        messages.error(request, e + " لم يتغيّر الجدول السابق.")
    if codes:
        messages.success(request, f"تم استيراد الجدول الزمني لـ {len(codes)} قسماً: {'، '.join(codes)}. "
                                  "حلّ محلّ الجدول السابق.")
    return redirect("timetable_home")


@student_required
def my_timetable(request):
    section = request.user.section
    return render(request, "quiz/timetable/mine.html", {
        "section": section, "grid": build_grid(section.code) if section else None,
    })


def student_grid(user):
    if not user.is_student or not user.section_id:
        return None
    return build_grid(user.section.code)

