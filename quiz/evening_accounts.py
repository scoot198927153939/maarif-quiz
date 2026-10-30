"""محاسبة الحصص المسائية: صفحة الدفع (مبلغ + رقم وصل + الأشهر التي يخصّها، يُقسم بالتساوي)
وصفحة العرض لكل قسم (التلميذ وأرقامه، والأشهر من منتصف أكتوبر حتى منتصف يوليو، ومجموع ما دفع).
للمدير (الأدمن) فقط مثل باقي الحصص المسائية."""
from urllib.parse import quote

from django.contrib import messages
from django.db.models import Q
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .attendance import _date
from .excel import HEADER_FILL, HEADER_FONT, _clean
from .models import (
    EVENING_PERIODS, EveningGroup, EveningPayment, EveningStudent, evening_period_dates, evening_school_year,
)
from .views import admin_required, xlsx_response

SCHOOL = "مدارس المعارف الحرة"


def _year(request):
    try:
        return int(request.GET.get("year") or request.POST.get("year"))
    except (TypeError, ValueError):
        return evening_school_year()


def _year_ctx(year):
    return {"year": year, "year_label": f"{year}-{year + 1}", "prev_year": year - 1, "next_year": year + 1}


def _periods(year):
    out = []
    for i, name in enumerate(EVENING_PERIODS):
        start, end = evening_period_dates(year, i)
        out.append({"index": i, "name": name, "start": start, "end": end,
                    "range": f"{start:%d/%m}–{end:%d/%m}"})
    return out


def _due_periods(student, periods, today):
    """الأشهر التي بدأت ولم يكن التلميذ قد سجّل بعد نهايتها."""
    reg = student.registered_on
    return [p["index"] for p in periods if p["start"] <= today and (reg is None or reg < p["end"])]


def _account_rows(group, year):
    """لكل تلميذ: المدفوع في كل شهر، المجموع، المطلوب حتى اليوم، المتبقي، والأشهر غير المسدّدة."""
    periods = _periods(year)
    today = timezone.localdate()
    fee = group.monthly_fee
    payments = list(EveningPayment.objects.filter(student__group=group, year=year))
    paid_ids = {p.student_id for p in payments}
    students = list(group.students.filter(Q(is_active=True) | Q(id__in=paid_ids)))
    by_student = {}
    for p in payments:
        cells = by_student.setdefault(p.student_id, [0] * len(periods))
        for i, share in p.shares.items():
            cells[i] += share
    rows = []
    for st in students:
        cells = by_student.get(st.id, [0] * len(periods))
        due = _due_periods(st, periods, today)
        unpaid = [i for i in due if fee and cells[i] < fee]
        expected = fee * len(due)
        total = sum(cells)
        rows.append({
            "student": st, "total": total, "expected": expected,
            "remaining": max(0, expected - total) if fee else None, "unpaid": unpaid,
            "unpaid_amount": sum(fee - cells[i] for i in unpaid),
            "cells": [{"amount": a, "state": ("full" if fee and a >= fee else "part" if a else
                                              "late" if (fee and i in due) else "")}
                      for i, a in enumerate(cells)],
        })
    col_totals = [sum(r["cells"][i]["amount"] for r in rows) for i in range(len(periods))]
    return periods, rows, col_totals


# ---------------------------------------------------------------- صفحة العرض

@admin_required
def evening_accounts(request):
    group = EveningGroup.objects.first()
    if group is None:
        messages.error(request, "أضف قسماً مسائياً أولاً.")
        return redirect("evening_home")
    return redirect("evening_account", group_id=group.id)


@admin_required
def evening_account(request, group_id):
    group = get_object_or_404(EveningGroup, pk=group_id)
    year = _year(request)
    if request.method == "POST":
        try:
            group.monthly_fee = max(0, int(_clean(request.POST.get("monthly_fee")) or 0))
        except ValueError:
            messages.error(request, "اكتب الرسوم الشهرية رقماً، مثل 1500.")
        else:
            group.save(update_fields=["monthly_fee"])
            messages.success(request, f"الرسوم الشهرية للقسم {group}: {group.monthly_fee} أوقية.")
        return redirect(f"{reverse('evening_account', args=[group.id])}?year={year}")
    periods, rows, col_totals = _account_rows(group, year)
    grand = sum(col_totals)
    if request.GET.get("export") == "xlsx":
        return xlsx_response(_account_workbook(group, year, periods, rows, col_totals),
                             f"accounts_{group.code}_{year}-{year + 1}.xlsx")
    return render(request, "quiz/evening/accounts.html", {
        "group": group, "groups": EveningGroup.objects.all(), "tab": "accounts", "periods": periods, "rows": rows,
        "col_totals": col_totals, "grand": grand, "n_late": sum(1 for r in rows if r["unpaid"]),
        "expected": sum(r["expected"] for r in rows), **_year_ctx(year),
    })


def _account_workbook(group, year, periods, rows, col_totals):
    wb = Workbook()
    ws = wb.active
    ws.title = group.code[:31]
    ws.sheet_view.rightToLeft = True
    ws.append([f"{SCHOOL} · محاسبة الحصص المسائية · {group.code} · {year}-{year + 1}"])
    head = ["#", "رقم القيد", "اسم التلميذ", "هاتف الوكيل", "الواتساب"]
    head += [f"{p['name']} ({p['range']})" for p in periods] + ["المجموع"]
    if group.monthly_fee:
        head += ["المطلوب حتى اليوم", "المتبقي"]
    ws.append(head)
    fills = {"full": PatternFill("solid", fgColor="D9F2E3"), "part": PatternFill("solid", fgColor="FFF1D6"),
             "late": PatternFill("solid", fgColor="FDE2DF")}
    for n, r in enumerate(rows, start=1):
        st = r["student"]
        line = [n, st.matricule, st.full_name, st.guardian_phone, st.whatsapp]
        line += [c["amount"] or None for c in r["cells"]] + [r["total"]]
        if group.monthly_fee:
            line += [r["expected"], r["remaining"]]
        ws.append(line)
        for i, c in enumerate(r["cells"]):
            if c["state"]:
                ws.cell(row=ws.max_row, column=6 + i).fill = fills[c["state"]]
    ws.append(["", "", "المجموع", "", ""] + col_totals + [sum(col_totals)])
    ws["A1"].font = Font(bold=True, size=13)
    for cell in ws[2]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for i, w in enumerate([5, 12, 30, 14, 14] + [13] * len(periods) + [12, 14, 12], start=1):
        ws.column_dimensions[ws.cell(row=2, column=i).column_letter].width = w
    return wb


@admin_required
def evening_late(request, group_id):
    """المتأخرون في الدفع: الأشهر التي بدأت ولم تُسدَّد كاملة، مع رسالة واتساب جاهزة للوكيل."""
    group = get_object_or_404(EveningGroup, pk=group_id)
    year = _year(request)
    periods, rows, _ = _account_rows(group, year)
    late = [r for r in rows if r["unpaid"]]
    for r in late:
        r["unpaid_names"] = "، ".join(EVENING_PERIODS[i] for i in r["unpaid"])
        st = r["student"]
        if st.wa_link:
            text = (f"السلام عليكم، نذكّركم بأن رسوم الحصص المسائية للتلميذ(ة) {st.full_name} ({group.code}) "
                    f"لم تُسدَّد كاملة عن: {r['unpaid_names']}. المبلغ المتبقي: {r['unpaid_amount']} أوقية. "
                    f"{SCHOOL}.")
            r["wa_message"] = f"{st.wa_link}?text={quote(text)}"
    return render(request, "quiz/evening/late.html", {
        "group": group, "groups": EveningGroup.objects.all(), "tab": "late", "rows": late,
        "total_late": sum(r["unpaid_amount"] for r in late), **_year_ctx(year),
    })


# ---------------------------------------------------------------- صفحة الدفع

@admin_required
def evening_pay(request, group_id):
    group = get_object_or_404(EveningGroup, pk=group_id)
    year = _year(request)
    students = list(group.students.filter(is_active=True))
    here = f"{reverse('evening_pay', args=[group.id])}?year={year}"
    if request.method == "POST":
        st = EveningStudent.objects.filter(group=group, pk=request.POST.get("student") or 0).first()
        receipt = _clean(request.POST.get("receipt_no"))[:40]
        months = sorted({int(m) for m in request.POST.getlist("months")
                         if m.isdigit() and int(m) < len(EVENING_PERIODS)})
        errors = []
        try:
            amount = int(_clean(request.POST.get("amount")).replace(" ", "") or 0)
        except ValueError:
            amount = 0
        if st is None:
            errors.append("اختر التلميذ.")
        if amount <= 0:
            errors.append("اكتب المبلغ المدفوع رقماً أكبر من 0.")
        if not receipt:
            errors.append("اكتب رقم الوصل.")
        elif EveningPayment.objects.filter(receipt_no__iexact=receipt, year=year).exists():
            errors.append(f"رقم الوصل {receipt} مسجّل من قبل. تأكد من الرقم.")
        if not months:
            errors.append("حدد الشهر أو الأشهر التي يخصّها الدفع.")
        if errors:
            for e in errors:
                messages.error(request, e)
            return render(request, "quiz/evening/pay.html", _pay_ctx(group, year, students, st, request.POST))
        p = EveningPayment.objects.create(
            student=st, year=year, amount=amount, receipt_no=receipt, months=",".join(map(str, months)),
            paid_on=_date(request.POST.get("paid_on")), note=_clean(request.POST.get("note"))[:255],
            recorded_by=request.user,
        )
        shares = p.shares
        detail = "، ".join(f"{EVENING_PERIODS[i]} {a}" for i, a in shares.items())
        messages.success(request, f"تم تسجيل وصل {receipt} للتلميذ {st.full_name}: {amount} أوقية ({detail}).")
        return redirect(f"{here}&student={st.id}&saved={p.id}")
    st = EveningStudent.objects.filter(group=group, pk=request.GET.get("student") or 0).first()
    return render(request, "quiz/evening/pay.html", _pay_ctx(group, year, students, st, None,
                                                            saved=request.GET.get("saved")))


def _pay_ctx(group, year, students, student, data, saved=None):
    periods = _periods(year)
    history = []
    if student:
        history = list(student.payments.filter(year=year))
        paid = [0] * len(periods)
        for p in history:
            for i, a in p.shares.items():
                paid[i] += a
        for per in periods:
            per["paid"] = paid[per["index"]]
    recent = (EveningPayment.objects.filter(student__group=group, year=year)
              .select_related("student").order_by("-created_at")[:20])
    picked = set(data.getlist("months")) if data else set()
    for per in periods:
        per["checked"] = str(per["index"]) in picked
    return {
        "group": group, "groups": EveningGroup.objects.all(), "tab": "pay", "students": students,
        "student": student, "periods": periods, "history": history, "recent": recent,
        "form": data or {}, "today": timezone.localdate(), "saved": saved,
        "history_total": sum(p.amount for p in history), **_year_ctx(year),
    }


@admin_required
@require_POST
def evening_payment_delete(request, pk):
    p = get_object_or_404(EveningPayment.objects.select_related("student"), pk=pk)
    st = p.student
    p.delete()
    messages.success(request, f"تم حذف وصل {p.receipt_no} ({p.amount} أوقية) للتلميذ {st.full_name}.")
    return redirect(f"{reverse('evening_pay', args=[st.group_id])}?year={p.year}&student={st.id}")


@admin_required
def evening_receipt(request, pk):
    p = get_object_or_404(EveningPayment.objects.select_related("student__group", "recorded_by"), pk=pk)
    periods = _periods(p.year)
    lines = [{"name": periods[i]["name"], "range": periods[i]["range"], "amount": a} for i, a in p.shares.items()]
    return render(request, "quiz/evening/receipt.html", {"p": p, "lines": lines, "school": SCHOOL,
                                                         **_year_ctx(p.year)})
