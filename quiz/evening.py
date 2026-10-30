"""الحصص المسائية: أقسام مستقلة (مثل G1_7MA)، تلاميذها من ملف Excel بتاريخ التسجيل،
وغياب يومي ولائحة الغائبين بأرقام وكلائهم. للمدير (الأدمن) فقط حالياً."""
import datetime
import re

from django.contrib import messages
from django.db import transaction
from django.db.models import Count, Q
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

from .attendance import _date, _day_label, _month, _month_nav
from .excel import HEADER_FILL, HEADER_FONT, _clean, _header_columns, _norm_header, _style_header, is_retired
from .forms import ImportFileForm
from .models import (
    ABSENT, ATTENDANCE_CHOICES, EVENING_HOURS, LATE, PRESENT, EveningAttendance, EveningGroup, EveningStudent,
    EveningTeacher, EveningTeacherAttendance, User,
)
from .views import admin_required, xlsx_response

STATUS_LABELS = dict(ATTENDANCE_CHOICES)
TEMPLATE_HEADERS = ["رقم القيد", "اسم التلميذ", "تاريخ التسجيل", "هاتف الوكيل", "الواتساب", "القسم"]
GROUP_RE = re.compile(r"^[A-Za-z0-9_\-]{1,30}$")


def _group_code(value):
    return _clean(value).upper().replace(" ", "").replace("-", "_")


def _parse_date(value):
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    text = _clean(value)
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%Y/%m/%d", "%d.%m.%Y"):
        try:
            return datetime.datetime.strptime(text.split(" ")[0], fmt).date()
        except ValueError:
            pass
    return None


def _columns(header):
    cols = _header_columns(header)
    cols.pop("section", None)
    for i, v in enumerate(header):
        h = _norm_header(v)
        if ("تاريخ" in h or "date" in h or "تسجيل" in h or "inscri" in h) and "date" not in cols:
            cols["date"] = i
        elif ("قسم" in h or "فصل" in h or "groupe" in h or "classe" in h) and "group" not in cols:
            cols["group"] = i
    if "name" in cols and cols["name"] == cols.get("date"):
        del cols["date"]
    return cols


def import_students(file, default_group):
    """يضيف تلاميذ إلى قسم مسائي أو يحدّث بياناتهم (برقم القيد). لا يُحذف أحد.
    إن كان في الملف عمود «القسم» باسم قسم مسائي موجود يُوضع التلميذ فيه.
    يعيد (المضافون، المحدَّثون، الأخطاء)."""
    try:
        wb = load_workbook(file, read_only=True, data_only=True)
    except Exception:
        return 0, 0, ["تعذّر قراءة الملف. تأكد أنه ملف Excel بصيغة ‎.xlsx"]
    groups = {g.code.upper(): g for g in EveningGroup.objects.all()}
    errors, parsed, seen = [], [], set()
    today = timezone.localdate()
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        header_idx = next((i for i, r in enumerate(rows[:6]) if len(_columns(r)) >= 2), None)
        if header_idx is None:
            if any(any(_clean(v) for v in r) for r in rows):
                errors.append(f"الورقة «{ws.title}»: لم أجد سطر العناوين (رقم القيد، اسم التلميذ، تاريخ التسجيل…).")
            continue
        cols = _columns(rows[header_idx])
        if "matricule" not in cols or "name" not in cols:
            errors.append(f"الورقة «{ws.title}»: يجب أن يكون فيها عمود «رقم القيد» وعمود «اسم التلميذ».")
            continue

        def get(row, key):
            i = cols.get(key)
            return row[i] if i is not None and i < len(row) else None

        for idx, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
            mat, name = _clean(get(row, "matricule")), _clean(get(row, "name"))
            if not mat and not name:
                continue
            if not mat:
                errors.append(f"السطر {idx}: رقم القيد فارغ ({name}).")
                continue
            group = default_group
            gcode = _group_code(get(row, "group"))
            if gcode:
                group = groups.get(gcode)
                if group is None:
                    errors.append(f"السطر {idx}: القسم «{gcode}» غير موجود في الحصص المسائية. أضفه أولاً.")
                    continue
            if (group.id, mat) in seen:
                errors.append(f"السطر {idx}: رقم القيد {mat} مكرر في الملف.")
                continue
            seen.add((group.id, mat))
            raw_date = get(row, "date")
            day = _parse_date(raw_date)
            if _clean(raw_date) and day is None:
                errors.append(f"السطر {idx}: تاريخ التسجيل «{_clean(raw_date)}» غير مفهوم. اكتبه مثل 2026-10-19 أو 19/10/2026.")
                continue
            phone = _clean(get(row, "phone1"))
            wa = _clean(get(row, "whatsapp")) or _clean(get(row, "phone2"))
            retired = is_retired(get(row, "retired")) if "retired" in cols else False
            parsed.append((group, mat, name, day or today, phone, wa, retired))
    if errors:
        return 0, 0, errors
    if not parsed:
        return 0, 0, ["لم يتم العثور على أي تلميذ في الملف."]

    created = updated = 0
    with transaction.atomic():
        for group, mat, name, day, phone, wa, retired in parsed:
            st = EveningStudent.objects.filter(group=group, matricule=mat).first()
            day_student = User.objects.filter(matricule=mat, role=User.STUDENT).first()
            if st is None:
                st = EveningStudent(group=group, matricule=mat)
                created += 1
            else:
                updated += 1
            st.full_name = name or st.full_name or (day_student.full_name if day_student else "") or mat
            st.registered_on = day
            # إن لم يكن الرقم في الملف نأخذه من لائحة تلاميذ المدرسة
            st.guardian_phone = phone or st.guardian_phone or (day_student.guardian_phone if day_student else "")
            st.whatsapp = wa or st.whatsapp or (day_student.wa_number if day_student else "")
            st.is_active = not retired
            st.save()
            st.ensure_account()
    return created, updated, []


def _students(group, day=None):
    qs = group.students.filter(is_active=True)
    if day:
        qs = qs.filter(Q(registered_on__isnull=True) | Q(registered_on__lte=day))
    return list(qs)


def _nav(day):
    return {"prev_day": (day - datetime.timedelta(days=1)).isoformat(),
            "next_day": (day + datetime.timedelta(days=1)).isoformat()}


# ---------------------------------------------------------------- الصفحة الرئيسية

@admin_required
def evening_home(request):
    if request.method == "POST":
        code = _group_code(request.POST.get("code"))
        if not GROUP_RE.match(code or ""):
            messages.error(request, "اكتب اسم القسم بالحروف اللاتينية والأرقام، مثل G2_7MA.")
        elif EveningGroup.objects.filter(code__iexact=code).exists():
            messages.error(request, f"القسم {code} موجود من قبل.")
        else:
            EveningGroup.objects.create(code=code)
            messages.success(request, f"تمت إضافة القسم {code}.")
        return redirect("evening_home")
    today = timezone.localdate()
    groups = EveningGroup.objects.annotate(
        n_students=Count("students", filter=Q(students__is_active=True), distinct=True),
        n_absent=Count("students__attendance", filter=Q(students__attendance__date=today,
                                                         students__attendance__status=ABSENT), distinct=True),
        taken=Count("students__attendance", filter=Q(students__attendance__date=today), distinct=True),
    )
    return render(request, "quiz/evening/home.html", {
        "groups": groups, "today": today, "today_label": _day_label(today),
    })


@admin_required
@require_POST
def evening_group_delete(request, group_id):
    group = get_object_or_404(EveningGroup, pk=group_id)
    if group.students.exists():
        messages.error(request, f"لا يمكن حذف القسم {group} لأن فيه تلاميذ.")
    else:
        group.delete()
        messages.success(request, f"تم حذف القسم {group}.")
    return redirect("evening_home")


# ---------------------------------------------------------------- الغياب

@admin_required
def evening_roll(request, group_id):
    group = get_object_or_404(EveningGroup, pk=group_id)
    day = _date(request.POST.get("date") or request.GET.get("date"))
    students = _students(group, day)
    here = f"{reverse('evening_roll', args=[group.id])}?date={day.isoformat()}"
    if request.method == "POST":
        if day > timezone.localdate():
            messages.error(request, "لا يمكن تسجيل الغياب ليوم لم يأتِ بعد.")
            return redirect(here)
        for st in students:
            status = request.POST.get(f"s_{st.id}")
            if status in STATUS_LABELS:
                EveningAttendance.objects.update_or_create(student=st, date=day, defaults={"status": status})
        messages.success(request, f"تم حفظ غياب {group} ليوم {_day_label(day)}.")
        return redirect(here)
    records = dict(EveningAttendance.objects.filter(student__in=students, date=day).values_list("student_id", "status"))
    rows = [{"student": st, "status": records.get(st.id, PRESENT)} for st in students]
    return render(request, "quiz/evening/roll.html", {
        "group": group, "tab": "roll", "rows": rows, "day": day, "day_label": _day_label(day),
        "taken": bool(records), "choices": ATTENDANCE_CHOICES, "is_future": day > timezone.localdate(),
        **_nav(day),
    })


@admin_required
def evening_absent(request, group_id):
    """الغائبون في يوم أو في فترة، مع هاتف الوكيل والواتساب."""
    group = get_object_or_404(EveningGroup, pk=group_id)
    start = _date(request.GET.get("from") or request.GET.get("date"))
    end = _date(request.GET.get("to"), start)
    if end < start:
        start, end = end, start
    records = (EveningAttendance.objects.filter(student__group=group, date__range=(start, end), status__in=[ABSENT, LATE])
               .select_related("student").order_by("student__matricule", "date"))
    by_student = {}
    for a in records:
        r = by_student.setdefault(a.student_id, {"student": a.student, "absent": [], "late": []})
        r["absent" if a.status == ABSENT else "late"].append(a.date)
    rows = [r for r in by_student.values() if r["absent"]]
    rows.sort(key=lambda r: (len(r["student"].matricule), r["student"].matricule))
    one_day = start == end
    title = _day_label(start) if one_day else f"من {start.isoformat()} إلى {end.isoformat()}"
    if request.GET.get("export") == "xlsx":
        return xlsx_response(_absent_workbook(group, rows, title, one_day),
                             f"absent_{group.code}_{start.isoformat()}" + ("" if one_day else f"_{end.isoformat()}") + ".xlsx")
    return render(request, "quiz/evening/absent.html", {
        "group": group, "tab": "absent", "rows": rows, "start": start, "end": end, "one_day": one_day,
        "title": title, "taken": EveningAttendance.objects.filter(student__group=group, date__range=(start, end)).exists(),
        **_nav(start),
    })


def _absent_workbook(group, rows, title, one_day):
    wb = Workbook()
    ws = wb.active
    ws.title = group.code[:31]
    ws.sheet_view.rightToLeft = True
    ws.append([f"الغائبون · الحصص المسائية · {group.code} · {title}"])
    head = ["#", "رقم القيد", "اسم التلميذ", "هاتف الوكيل", "الواتساب"]
    if not one_day:
        head += ["عدد الغيابات", "أيام الغياب"]
    ws.append(head)
    for i, r in enumerate(rows, start=1):
        st = r["student"]
        line = [i, st.matricule, st.full_name, st.guardian_phone, st.whatsapp]
        if not one_day:
            line += [len(r["absent"]), "، ".join(d.isoformat() for d in r["absent"])]
        ws.append(line)
    ws["A1"].font = Font(bold=True, size=13)
    for cell in ws[2]:
        cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, Alignment(horizontal="center")
    for col, w in zip("ABCDEFG", [6, 14, 32, 16, 16, 12, 40]):
        ws.column_dimensions[col].width = w
    return wb


# ---------------------------------------------------------------- التلاميذ

@admin_required
def evening_students(request, group_id):
    group = get_object_or_404(EveningGroup, pk=group_id)
    if request.method == "POST":
        form = ImportFileForm(request.POST, request.FILES)
        if form.is_valid():
            created, updated, errors = import_students(form.cleaned_data["file"], group)
            if errors:
                for e in errors[:15]:
                    messages.error(request, e)
                if len(errors) > 15:
                    messages.error(request, f"… و{len(errors) - 15} خطأ آخر.")
                messages.error(request, "لم يُستورد شيء. صحّح الملف ثم أعد المحاولة.")
            else:
                messages.success(request, f"تم الاستيراد: {created} تلميذ جديد، و{updated} تلميذ حُدّثت بياناته.")
        else:
            messages.error(request, "اختر ملف Excel أولاً.")
        return redirect("evening_students", group_id=group.id)
    students = group.students.filter(is_active=True).select_related("user").annotate(
        n_absent=Count("attendance", filter=Q(attendance__status=ABSENT)))
    return render(request, "quiz/evening/students.html", {
        "group": group, "tab": "students", "students": students, "import_form": ImportFileForm(),
    })


@admin_required
@require_POST
def evening_import(request):
    """استيراد اللائحة من الصفحة الرئيسية: يُختار القسم من القائمة."""
    group = get_object_or_404(EveningGroup, pk=request.POST.get("group") or 0)
    return evening_students(request, group.id)


@admin_required
@require_POST
def evening_student_add(request, group_id):
    """إضافة تلميذ واحد يدوياً بدون Excel (أو تحديث بياناته إن كان رقم قيده موجوداً)."""
    group = get_object_or_404(EveningGroup, pk=group_id)
    mat = _clean(request.POST.get("matricule"))
    name = _clean(request.POST.get("full_name"))
    day = _parse_date(request.POST.get("registered_on")) or timezone.localdate()
    if not mat:
        messages.error(request, "اكتب رقم القيد.")
        return redirect("evening_students", group_id=group.id)
    day_student = User.objects.filter(matricule=mat, role=User.STUDENT).first()
    name = name or (day_student.full_name if day_student else "")
    if not name:
        messages.error(request, "اكتب اسم التلميذ.")
        return redirect("evening_students", group_id=group.id)
    st, created = EveningStudent.objects.get_or_create(group=group, matricule=mat, defaults={"full_name": name})
    st.full_name = name
    st.registered_on = day
    st.guardian_phone = (_clean(request.POST.get("guardian_phone")) or st.guardian_phone
                         or (day_student.guardian_phone if day_student else ""))
    st.whatsapp = (_clean(request.POST.get("whatsapp")) or st.whatsapp
                   or (day_student.wa_number if day_student else ""))
    st.is_active = True
    st.save()
    st.ensure_account()
    messages.success(request, f"{'تمت إضافة' if created else 'تم تحديث'} {st.full_name} في القسم {group}.")
    return redirect("evening_students", group_id=group.id)


@admin_required
@require_POST
def evening_student_remove(request, es_id):
    st = get_object_or_404(EveningStudent, pk=es_id)
    st.is_active = False
    st.save(update_fields=["is_active"])
    messages.success(request, f"تم حذف {st.full_name} من القسم {st.group}. (يبقى غيابه السابق محفوظاً، ويعود إذا استوردته مرة أخرى.)")
    return redirect("evening_students", group_id=st.group_id)


@admin_required
def evening_template(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "التلاميذ"
    ws.append(TEMPLATE_HEADERS)
    ws.append(["12345", "محمد الأمين أحمد", "2026-10-19", "22223344", "36123456", "G1_7MA"])
    ws.append(["12346", "فاطمة سيدي", "2026-10-20", "44556677", "", ""])
    _style_header(ws, [14, 30, 16, 16, 16, 12])
    return xlsx_response(wb, "evening_students_template.xlsx")


# ---------------------------------------------------------------- حضور الأساتذة

@admin_required
def evening_teachers(request):
    """حضور أساتذة الحصص المسائية اليومي: لكل أستاذ مجموعة واحدة وأربع ساعات متتالية."""
    day = _date(request.POST.get("date") or request.GET.get("date"))
    here = f"{reverse('evening_teachers')}?date={day.isoformat()}"
    teachers = list(EveningTeacher.objects.filter(is_active=True).select_related("group"))
    groups = {str(g.id): g for g in EveningGroup.objects.all()}
    if request.method == "POST":
        if "add_teacher" in request.POST:
            name = " ".join(request.POST.get("name", "").split())[:150]
            if name:
                t, created = EveningTeacher.objects.get_or_create(name=name)
                t.is_active = True
                t.group = groups.get(request.POST.get("group", "")) or t.group
                t.save()
                messages.success(request, f"تمت إضافة الأستاذ «{t}»." if created else f"الأستاذ «{t}» موجود في اللائحة.")
            return redirect(here)
        if "remove_teacher" in request.POST:
            EveningTeacher.objects.filter(pk=request.POST.get("remove_teacher")).update(is_active=False)
            messages.success(request, "تم حذف الأستاذ من اللائحة (يبقى سجل حضوره).")
            return redirect(here)
        if day > timezone.localdate():
            messages.error(request, "لا يمكن تسجيل الحضور ليوم لم يأتِ بعد.")
            return redirect(here)
        saved = 0
        for t in teachers:
            status = request.POST.get(f"s_{t.id}", "")
            group = groups.get(request.POST.get(f"g_{t.id}", ""))
            if status in STATUS_LABELS:
                try:
                    hours = max(0, min(12, int(request.POST.get(f"h_{t.id}") or EVENING_HOURS)))
                except ValueError:
                    hours = EVENING_HOURS
                if status == ABSENT:
                    hours = 0
                EveningTeacherAttendance.objects.update_or_create(
                    teacher=t, date=day,
                    defaults={"status": status, "group": group, "hours": hours,
                              "note": request.POST.get(f"n_{t.id}", "").strip()[:255]},
                )
                saved += 1
            else:  # «لا حصة»
                EveningTeacherAttendance.objects.filter(teacher=t, date=day).delete()
            if group and t.group_id != group.id:  # نتذكّر مجموعة الأستاذ للأيام القادمة
                t.group = group
                t.save(update_fields=["group"])
        messages.success(request, f"تم حفظ حضور {saved} أستاذ ليوم {_day_label(day)}.")
        return redirect(here)

    records = {a.teacher_id: a for a in EveningTeacherAttendance.objects.filter(date=day)}
    rows = []
    for t in teachers:
        rec = records.get(t.id)
        rows.append({"teacher": t, "status": rec.status if rec else "", "note": rec.note if rec else "",
                     "group": (rec.group_id if rec else None) or t.group_id,
                     "hours": rec.hours if rec else EVENING_HOURS})
    return render(request, "quiz/evening/teachers.html", {
        "rows": rows, "groups": list(groups.values()), "day": day, "day_label": _day_label(day),
        "taken": bool(records), "choices": ATTENDANCE_CHOICES, "is_future": day > timezone.localdate(),
        "month": day.strftime("%Y-%m"), **_nav(day),
    })


@admin_required
def evening_teachers_monthly(request):
    """الحضور والغياب الشهري لكل أستاذ: أيام الحضور والغياب والتأخر ومجموع الساعات."""
    start, end = _month(request.GET.get("month"))
    records = list(EveningTeacherAttendance.objects.filter(date__gte=start, date__lt=end)
                   .select_related("group").order_by("date"))
    stats = {}
    for a in records:
        st = stats.setdefault(a.teacher_id, {"P": 0, "A": 0, "L": 0, "hours": 0, "absent_days": [], "groups": set()})
        st[a.status] += 1
        st["hours"] += a.hours
        if a.status == ABSENT:
            st["absent_days"].append(a.date)
        if a.group:
            st["groups"].add(a.group.code)
    teachers = [t for t in EveningTeacher.objects.select_related("group") if t.is_active or t.id in stats]
    rows = []
    for t in teachers:
        st = stats.get(t.id, {"P": 0, "A": 0, "L": 0, "hours": 0, "absent_days": [], "groups": set()})
        rows.append({"teacher": t, "present": st["P"], "absent": st["A"], "late": st["L"], "hours": st["hours"],
                     "absent_days": st["absent_days"],
                     "groups": sorted(st["groups"]) or ([t.group.code] if t.group else [])})
    n_days = len({a.date for a in records})
    nav = _month_nav(start)
    if request.GET.get("export") == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = start.strftime("%Y-%m")
        ws.sheet_view.rightToLeft = True
        ws.append([f"حضور أساتذة الحصص المسائية · {nav['month_label']} · {n_days} يوم مسجّل"])
        ws.append(["#", "الأستاذ", "المجموعة", "حاضر", "متأخر", "غائب", "مجموع الساعات", "أيام الغياب"])
        for i, r in enumerate(rows, start=1):
            ws.append([i, r["teacher"].name, "، ".join(r["groups"]), r["present"], r["late"], r["absent"], r["hours"],
                       "، ".join(d.strftime("%d/%m") for d in r["absent_days"])])
        ws["A1"].font = Font(bold=True, size=13)
        for cell in ws[2]:
            cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, Alignment(horizontal="center")
        for col, w in zip("ABCDEFGH", [5, 26, 14, 8, 8, 8, 14, 30]):
            ws.column_dimensions[col].width = w
        return xlsx_response(wb, f"evening_teachers_{nav['month']}.xlsx")
    return render(request, "quiz/evening/teachers_monthly.html", {"rows": rows, "n_days": n_days, **nav})
