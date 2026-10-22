"""سجل النشاط: من متصل الآن، ومن أضاف أو عدّل أو حذف ماذا، مع تقرير يومي وشهري."""
import datetime

from django.apps import apps
from django.contrib.auth.signals import user_logged_in, user_logged_out, user_login_failed
from django.core.paginator import Paginator
from django.db.models import Count, Max, Min, Q
from django.db.models.functions import TruncDate
from django.dispatch import receiver
from django.shortcuts import render
from django.urls import Resolver404, resolve
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font

from .attendance import MONTHS, _date, _day_label, _month, _month_nav
from .models import (
    ActivityLog, Attempt, Classe, EveningGroup, Exam, Question, Section, StudentNote, Teacher, User,
)
from .views import admin_required, xlsx_response

# من كان له طلب خلال هذه المدة يُعدّ متصلاً
ONLINE_MINUTES = 5
# لا نكتب «آخر نشاط» في قاعدة البيانات أكثر من مرة كل دقيقة لكل مستخدم
SEEN_EVERY = datetime.timedelta(seconds=60)

A = ActivityLog
# اسم الصفحة (url name) ← (نوع الحدث، وصفه). تُسجَّل طلبات POST الناجحة فقط.
ACTIONS = {
    "password_change": (A.EDIT, "تغيير كلمة المرور"),
    "exam_create": (A.ADD, "إنشاء امتحان"),
    "exam_edit": (A.EDIT, "تعديل امتحان"),
    "exam_delete": (A.DELETE, "حذف امتحان"),
    "exam_import": (A.IMPORT, "استيراد أسئلة امتحان من Excel"),
    "question_add": (A.ADD, "إضافة سؤال"),
    "question_edit": (A.EDIT, "تعديل سؤال"),
    "question_delete": (A.DELETE, "حذف سؤال"),
    "attempt_reset": (A.EDIT, "السماح بإعادة امتحان"),
    "user_create": (A.ADD, "إضافة مستخدم"),
    "user_edit": (A.EDIT, "تعديل مستخدم"),
    "user_delete": (A.DELETE, "حذف مستخدم"),
    "students_import": (A.IMPORT, "استيراد التلاميذ من Excel"),
    "roster_import": (A.IMPORT, "استيراد لائحة التلاميذ"),
    "teachers_import": (A.IMPORT, "استيراد لائحة الأساتذة"),
    "timetable_import": (A.IMPORT, "استيراد الجدول الزمني"),
    "timetable_edit": (A.EDIT, "تعديل الجدول الزمني"),
    "lesson_new": (A.ADD, "تصوير درس"),
    "video_new": (A.ADD, "إضافة فيديو تعليمي"),
    "video_edit": (A.EDIT, "تعديل فيديو تعليمي"),
    "video_delete": (A.DELETE, "حذف فيديو تعليمي"),
    "lesson_delete": (A.DELETE, "حذف درس مصوّر"),
    "lesson_scopes": (A.EDIT, "تحديد مواد مسجّل الدروس"),
    "section_roll": (A.EDIT, "تسجيل حضور التلاميذ"),
    "evening_home": (A.ADD, "إضافة قسم للحصص المسائية"),
    "evening_group_delete": (A.DELETE, "حذف قسم من الحصص المسائية"),
    "evening_roll": (A.EDIT, "تسجيل غياب الحصص المسائية"),
    "evening_students": (A.IMPORT, "استيراد تلاميذ الحصص المسائية"),
    "evening_teachers": (A.EDIT, "تسجيل حضور أساتذة الحصص المسائية"),
    "evening_student_add": (A.ADD, "إضافة تلميذ للحصص المسائية"),
    "evening_import": (A.IMPORT, "استيراد تلاميذ الحصص المسائية"),
    "evening_student_remove": (A.DELETE, "حذف تلميذ من الحصص المسائية"),
    "student_file": (A.ADD, "إضافة ملاحظة على تلميذ"),
    "note_delete": (A.DELETE, "حذف ملاحظة"),
    "teacher_roll": (A.EDIT, "تسجيل حضور الأساتذة"),
    "teacher_file": (A.EDIT, "تعديل بيانات أستاذ"),
    "library_level": (A.EDIT, "تعديل مخزون الكتب"),
    "library_section": (A.EDIT, "تسليم الكتب"),
    "exam_intro": (A.EXAM, "بدء امتحان"),
    "submit_exam": (A.EXAM, "تسليم امتحان"),
    "report_violation": (A.EXAM, "خروج من شاشة الامتحان (مخالفة)"),
}
# لا تُسجَّل: الدخول والخروج لهما إشارات خاصة، وحفظ كل جواب كثير جداً
SKIP = {"login", "logout", "save_answer"}


def _target(kwargs):
    """وصف قصير للشيء الذي جرى عليه الحدث، يُقرأ قبل تنفيذ الطلب (حتى لا يضيع عند الحذف)."""
    try:
        if "exam_id" in kwargs:
            e = Exam.objects.select_related("classe").get(pk=kwargs["exam_id"])
            return f"{e.title} ({e.classe})"
        if "question_id" in kwargs:
            q = Question.all_objects.select_related("exam").get(pk=kwargs["question_id"])
            return f"سؤال في: {q.exam.title}"
        if "attempt_id" in kwargs:
            a = Attempt.objects.select_related("student", "exam").get(pk=kwargs["attempt_id"])
            return f"{a.student} — {a.exam.title}"
        for key in ("user_id", "student_id"):
            if key in kwargs:
                return str(User.objects.get(pk=kwargs[key]))
        if "section_id" in kwargs:
            return Section.objects.get(pk=kwargs["section_id"]).code
        if "classe_id" in kwargs:
            return Classe.objects.get(pk=kwargs["classe_id"]).code
        if "teacher_id" in kwargs:
            return Teacher.objects.get(pk=kwargs["teacher_id"]).name
        if "group_id" in kwargs:
            return EveningGroup.objects.get(pk=kwargs["group_id"]).code
        if "code" in kwargs:
            return kwargs["code"]
        if "video_id" in kwargs:
            from .models import Video
            v = Video.objects.select_related("classe").get(pk=kwargs["video_id"])
            return f"{v.display_title} ({v.classe})"
        if "note_id" in kwargs:
            return str(StudentNote.objects.select_related("student").get(pk=kwargs["note_id"]).student)
    except Exception:  # noqa: BLE001 — الوصف ثانوي، لا يجب أن يوقف الطلب
        pass
    return ""


def _admin_event(match, request):
    """تغيير من صفحة «الإعدادات» (لوحة Django)."""
    name = match.url_name or ""
    parts = name.split("_")
    if len(parts) < 3:
        return None
    op = parts[-1]
    model = None
    try:
        model = apps.get_model(parts[0], "_".join(parts[1:-1]))
    except LookupError:
        pass
    label = model._meta.verbose_name if model else "_".join(parts[1:-1])
    target = ""
    if model is not None and match.kwargs.get("object_id"):
        obj = model._default_manager.filter(pk=match.kwargs["object_id"]).first()
        target = str(obj) if obj else ""
    if op == "add":
        return A.ADD, f"الإعدادات: إضافة {label}", target
    if op == "change":
        return A.EDIT, f"الإعدادات: تعديل {label}", target
    if op == "delete":
        return A.DELETE, f"الإعدادات: حذف {label}", target
    if op == "changelist" and request.POST.get("action") == "delete_selected":
        n = len(request.POST.getlist("_selected_action"))
        return A.DELETE, f"الإعدادات: حذف {n} من {label}", ""
    return None


def describe(request):
    """يحدد الحدث الذي يمثّله طلب POST، أو None إن لم يكن يستحق التسجيل."""
    try:
        match = resolve(request.path_info)
    except Resolver404:
        return None
    name = match.url_name
    if name in SKIP:
        return None
    if match.namespace == "admin":
        event = _admin_event(match, request)
        if event is None:
            return None
        kind, action, target = event
    else:
        kind, action = ACTIONS.get(name, (A.OTHER, name or request.path))
        if name == "exam_toggle_publish":
            exam = Exam.objects.filter(pk=match.kwargs.get("exam_id")).first()
            kind, action = A.EDIT, ("إخفاء امتحان" if exam and exam.is_published else "نشر امتحان")
        target = _target(match.kwargs)
    return {"kind": kind, "action": action, "target": target[:255]}


def client_ip(request):
    forwarded = request.META.get("HTTP_X_FORWARDED_FOR", "")
    ip = forwarded.split(",")[0].strip() if forwarded else request.META.get("REMOTE_ADDR")
    return ip or None


def record(request, user, kind, action, target="", username=None):
    ActivityLog.objects.create(
        user=user if user is not None and user.pk else None,
        username=(str(user) if user is not None and user.pk else username or "")[:150],
        role=getattr(user, "effective_role", "") or "",
        kind=kind, action=action[:200], target=(target or "")[:255],
        path=request.path[:255] if request else "", ip=client_ip(request) if request else None,
    )


def touch(user):
    """يحدّث «آخر نشاط» للمستخدم (مرة في الدقيقة على الأكثر)."""
    now = timezone.now()
    if user.last_seen is None or now - user.last_seen > SEEN_EVERY:
        User.objects.filter(pk=user.pk).update(last_seen=now)
        user.last_seen = now


@receiver(user_logged_in)
def _on_login(sender, request, user, **kwargs):
    touch(user)
    record(request, user, A.LOGIN, "تسجيل الدخول")


@receiver(user_logged_out)
def _on_logout(sender, request, user, **kwargs):
    if user is not None:
        # لم يعد متصلاً
        User.objects.filter(pk=user.pk).update(last_seen=None)
        record(request, user, A.LOGOUT, "تسجيل الخروج")


@receiver(user_login_failed)
def _on_login_failed(sender, credentials, request=None, **kwargs):
    username = (credentials or {}).get("username", "")
    user = User.objects.filter(username=username).first() if username else None
    record(request, user, A.FAILED, "محاولة دخول بكلمة سر خاطئة", username=username or "?")


# ---------------------------------------------------------------- الصفحة

def _period(request):
    """يوم أو شهر، حسب اختيار المشرف. يعيد (بداية، نهاية، معلومات العرض)."""
    mode = "month" if request.GET.get("mode") == "month" else "day"
    if mode == "month":
        start, end = _month(request.GET.get("month"))
        info = {"mode": mode, "label": _month_nav(start)["month_label"], **_month_nav(start)}
    else:
        start = _date(request.GET.get("day"))
        end = start + datetime.timedelta(days=1)
        info = {
            "mode": mode, "label": _day_label(start), "day": start.isoformat(),
            "prev_day": (start - datetime.timedelta(days=1)).isoformat(),
            "next_day": end.isoformat(), "month": start.strftime("%Y-%m"),
        }
    tz = timezone.get_current_timezone()
    bounds = (
        timezone.make_aware(datetime.datetime.combine(start, datetime.time.min), tz),
        timezone.make_aware(datetime.datetime.combine(end, datetime.time.min), tz),
    )
    return bounds, info


ROLE_LABELS = dict(User.ROLE_CHOICES)
KIND_LABELS = dict(ActivityLog.KIND_CHOICES)


@admin_required
def activity_report(request):
    now = timezone.now()
    online = list(
        User.objects.filter(last_seen__gte=now - datetime.timedelta(minutes=ONLINE_MINUTES))
        .select_related("classe", "section").order_by("-is_superuser", "role", "-last_seen")
    )
    online_by_role = {}
    for u in online:
        online_by_role[u.effective_role_label] = online_by_role.get(u.effective_role_label, 0) + 1

    (since, until), info = _period(request)
    logs = ActivityLog.objects.filter(created_at__gte=since, created_at__lt=until)
    role = request.GET.get("role", "")
    kind = request.GET.get("kind", "")
    q = request.GET.get("q", "").strip()
    if role:
        logs = logs.filter(role=role)
    if kind:
        logs = logs.filter(kind=kind)
    if q:
        logs = logs.filter(Q(username__icontains=q) | Q(target__icontains=q) | Q(action__icontains=q))

    if request.GET.get("export"):
        return _export(logs, info)

    kinds = [
        {"kind": r["kind"], "label": KIND_LABELS.get(r["kind"], r["kind"]), "n": r["n"]}
        for r in logs.values("kind").annotate(n=Count("id")).order_by("-n")
    ]
    per_user = list(
        logs.values("username", "role").annotate(n=Count("id"), first=Min("created_at"), last=Max("created_at"))
        .order_by("-n")
    )
    for r in per_user:
        r["role_label"] = ROLE_LABELS.get(r["role"], "")
    per_day = []
    if info["mode"] == "month":
        per_day = list(
            logs.annotate(d=TruncDate("created_at")).values("d")
            .annotate(n=Count("id"), users=Count("username", distinct=True)).order_by("d")
        )
        for r in per_day:
            r["label"] = _day_label(r["d"])
            r["iso"] = r["d"].isoformat()
    page = Paginator(logs.select_related("user"), 200).get_page(request.GET.get("page"))
    for e in page:
        e.role_label = ROLE_LABELS.get(e.role, "")

    params = request.GET.copy()
    params.pop("page", None)
    params.pop("export", None)
    return render(request, "quiz/activity.html", {
        "online": online, "online_by_role": online_by_role, "online_minutes": ONLINE_MINUTES,
        "info": info, "page": page, "total": page.paginator.count, "kinds": kinds,
        "per_user": per_user, "per_day": per_day,
        "role": role, "kind": kind, "q": q, "query": params.urlencode(),
        "role_choices": User.ROLE_CHOICES, "kind_choices": ActivityLog.KIND_CHOICES,
    })


def _export(logs, info):
    wb = Workbook()
    ws = wb.active
    ws.title = "سجل النشاط"
    ws.sheet_view.rightToLeft = True
    ws.append([f"سجل النشاط — {info['label']}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["الوقت", "المستخدم", "الدور", "النوع", "العملية", "على", "IP"])
    for c in ws[2]:
        c.font = Font(bold=True)
    for e in logs.order_by("created_at"):
        ws.append([
            timezone.localtime(e.created_at).strftime("%Y-%m-%d %H:%M"), e.username,
            ROLE_LABELS.get(e.role, ""), KIND_LABELS.get(e.kind, e.kind), e.action, e.target, e.ip or "",
        ])
    for col, width in zip("ABCDEFG", (18, 28, 14, 14, 34, 40, 16)):
        ws.column_dimensions[col].width = width
    suffix = info.get("day") or info.get("month")
    return xlsx_response(wb, f"activity-{suffix}.xlsx")


__all__ = ["activity_report", "describe", "record", "touch", "MONTHS"]
